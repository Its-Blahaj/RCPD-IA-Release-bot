from __future__ import annotations

import json
import logging
import os
from pathlib import Path
from typing import Any

import discord
from discord import app_commands
from discord.ext import commands
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
CONFIG_DIR = ROOT / "config"
WORKBOOK_PATH = ROOT / "violations.xlsx"
load_dotenv(ROOT / ".env")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)
log = logging.getLogger("rcpd-ia")
TOKEN = os.getenv("DISCORD_TOKEN")
GUILD_ID = os.getenv("GUILD_ID", "").strip()


def load_json(filename: str) -> dict[str, Any]:
    with (CONFIG_DIR / filename).open("r", encoding="utf-8") as file:
        return json.load(file)


def load_config() -> dict[str, Any]:
    return load_json("config.json")


def load_templates() -> dict[str, Any]:
    return load_json("templates.json")


def load_options(filename: str) -> dict[str, Any]:
    return load_json(filename)


def format_list(items: list[str]) -> str:
    if len(items) == 0:
        return ""
    if len(items) == 1:
        return items[0]
    if len(items) == 2:
        return f"{items[0]} and {items[1]}"
    return f"{', '.join(items[:-1])}, and {items[-1]}"


def safe_text(value: str, max_length: int = 1000) -> str:
    value = value.strip()
    if not value:
        raise ValueError("This field cannot be empty.")
    if len(value) > max_length:
        raise ValueError(f"Please keep this field at or below {max_length} characters.")
    return value


def display_officer(user: discord.User | discord.Member | None, manual_name: str | None) -> str:
    # A manually entered name always wins, even if a Discord user is also selected.
    if manual_name and manual_name.strip():
        return discord.utils.escape_mentions(manual_name.strip())
    if user is not None:
        return user.mention
    raise ValueError("Select a Discord user or enter an officer name.")


def load_violation_categories() -> list[dict[str, Any]]:
    """Read every worksheet with ID/Offense/Severity/Punishment columns.

    Each offense worksheet is a category. Procedure/rule worksheets without an
    Offense column are deliberately excluded from the violation selector.
    """
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(
            f"Missing {WORKBOOK_PATH.name}. Put the Excel violation matrix next to bot.py."
        )

    workbook = load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    categories: list[dict[str, Any]] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            rows = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = [str(value).strip().casefold() if value is not None else "" for value in header_row]
            positions = {header: index for index, header in enumerate(headers) if header}
            if not all(key in positions for key in ("id", "offense", "severity", "punishment")):
                continue

            records: list[dict[str, str]] = []
            for row_number, row in enumerate(rows, start=2):
                def cell(name: str) -> str:
                    index = positions[name]
                    value = row[index] if index < len(row) else None
                    return str(value).strip() if value is not None else ""

                offense_id = cell("id")
                offense = cell("offense")
                if not offense_id or not offense:
                    continue
                records.append({
                    "key": f"{sheet_index}:{row_number}",
                    "id": offense_id,
                    "offense": offense,
                    "severity": cell("severity"),
                    "punishment": cell("punishment"),
                })
            if records:
                categories.append({
                    "key": str(sheet_index),
                    "name": sheet.title.strip(),
                    "records": records,
                })
    finally:
        workbook.close()

    if not categories:
        raise ValueError(
            "No offense worksheets found. Expected columns named ID, Offense, Severity, and Punishment."
        )
    return categories


def all_violation_records(categories: list[dict[str, Any]]) -> dict[str, dict[str, str]]:
    return {record["key"]: record for category in categories for record in category["records"]}


def authorized(interaction: discord.Interaction, config: dict[str, Any]) -> bool:
    if interaction.guild is None or not isinstance(interaction.user, discord.Member):
        return False
    member = interaction.user
    if member.guild_permissions.administrator:
        return True
    role_ids = {int(value) for value in config.get("authorized_role_ids", [])}
    return any(role.id in role_ids for role in member.roles)


def _format_template(value: str, values: dict[str, str]) -> str:
    try:
        return str(value).format_map(values)
    except KeyError as exc:
        raise ValueError(f"Template contains an unknown placeholder: {exc}") from exc


def render_template(release_type: str, values: dict[str, str]) -> str:
    template_data = load_templates().get(release_type)
    if not template_data:
        raise ValueError(f"Missing template for release type '{release_type}'.")
    embed_data = template_data.get("embed", {})
    description = embed_data.get("description", template_data.get("template", ""))
    return _format_template(description, values)


def build_release_embed(release_type: str, values: dict[str, str], config: dict[str, Any]) -> discord.Embed:
    templates = load_templates()
    template_data = templates.get(release_type)
    if not template_data:
        raise ValueError(f"Missing template for release type '{release_type}'.")
    data = template_data.get("embed", {})
    description_template = data.get("description", template_data.get("template", ""))
    title = _format_template(data.get("title", ""), values).strip()
    description = _format_template(description_template, values).strip()
    color_value = str(data.get("color") or config.get("embed_color", "#202225")).strip().lstrip("#")
    try:
        color = discord.Color(int(color_value, 16))
    except (TypeError, ValueError):
        color = discord.Color.dark_grey()
    embed = discord.Embed(title=title or None, description=description or None, color=color)
    url = _format_template(str(data.get("url", "")), values).strip()
    if url:
        embed.url = url
    author = _format_template(str(data.get("author", "")), values).strip()
    author_icon = _format_template(str(data.get("author_icon_url", "")), values).strip()
    if author:
        embed.set_author(name=author[:256], icon_url=author_icon or None)
    footer = _format_template(str(data.get("footer", config.get("embed_footer", ""))), values).strip()
    footer_icon = _format_template(str(data.get("footer_icon_url", "")), values).strip()
    if footer:
        embed.set_footer(text=footer[:2048], icon_url=footer_icon or None)
    thumbnail = _format_template(str(data.get("thumbnail", "")), values).strip()
    image = _format_template(str(data.get("image", "")), values).strip()
    if thumbnail:
        embed.set_thumbnail(url=thumbnail)
    if image:
        embed.set_image(url=image)
    fields = data.get("fields", [])
    if not isinstance(fields, list):
        raise ValueError("Template embed fields must be a JSON list.")
    if len(fields) > 25:
        raise ValueError("An embed can contain at most 25 fields.")
    for field in fields:
        if not isinstance(field, dict) or not str(field.get("name", "")).strip():
            continue
        field_name = _format_template(str(field.get("name", "")), values)
        field_value = _format_template(str(field.get("value", "")), values)
        if not field_value.strip():
            field_value = "—"
        embed.add_field(name=field_name[:256], value=field_value[:1024], inline=bool(field.get("inline", False)))
    return embed


def save_json(filename: str, data: dict[str, Any]) -> None:
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    target = CONFIG_DIR / filename
    temporary = target.with_suffix(target.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)
        file.write("\n")
    temporary.replace(target)


def release_embed(text: str, config: dict[str, Any]) -> discord.Embed:
    # Compatibility helper for older callers.
    color_value = str(config.get("embed_color", "#202225")).lstrip("#")
    try:
        color = discord.Color(int(color_value, 16))
    except (TypeError, ValueError):
        color = discord.Color.dark_grey()
    embed = discord.Embed(description=text, color=color)
    footer = str(config.get("embed_footer", "")).strip()
    if footer:
        embed.set_footer(text=footer)
    return embed


async def publish_preview(
    interaction: discord.Interaction,
    config: dict[str, Any],
    release_type: str,
    values: dict[str, str],
    author_id: int,
) -> None:
    if not authorized(interaction, config):
        await interaction.response.send_message(
            "You do not have an authorized IA role to create releases.", ephemeral=True
        )
        return
    try:
        embed = build_release_embed(release_type, values, config)
    except (ValueError, OSError, json.JSONDecodeError) as exc:
        await interaction.response.send_message(f"Template error: {exc}", ephemeral=True)
        return

    max_length = int(config.get("max_release_length", 4000))
    rendered_length = sum(len(str(part or "")) for part in [embed.title, embed.description, embed.footer.text if embed.footer else "", embed.author.name if embed.author else ""])
    rendered_length += sum(len(field.name) + len(field.value) for field in embed.fields)
    if rendered_length > min(max_length, 6000):
        await interaction.response.send_message(
            f"The rendered embed is too long ({rendered_length} characters; limit {min(max_length, 6000)}).",
            ephemeral=True,
        )
        return
    await interaction.response.send_message(
        content="**Preview — check the release before publishing:**",
        embed=embed,
        view=PublishView(embed.to_dict(), config, author_id),
        ephemeral=True,
        allowed_mentions=discord.AllowedMentions.none(),
    )


class PublishView(discord.ui.View):
    def __init__(self, embed_data: dict[str, Any], config: dict[str, Any], author_id: int,
                 template_key: str = "standard", values: dict[str, str] | None = None):
        super().__init__(timeout=300)
        self.embed_data = embed_data
        self.config = config
        self.author_id = author_id
        self.template_key = template_key
        self.values = values or {}

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who created this preview can use it.", ephemeral=True)
            return False
        if not authorized(interaction, self.config):
            await interaction.response.send_message("You no longer have permission to publish releases.", ephemeral=True)
            return False
        return True

    @discord.ui.button(label="Publish Release", style=discord.ButtonStyle.success, row=0)
    async def publish(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        channel_id = self.config.get("release_channel_id")
        channel = interaction.guild.get_channel(int(channel_id)) if interaction.guild and channel_id else None
        if not isinstance(channel, (discord.TextChannel, discord.Thread)):
            await interaction.response.send_message(
                "The release channel is not configured correctly in config/config.json.", ephemeral=True
            )
            return
        try:
            message = await channel.send(
                embed=discord.Embed.from_dict(self.embed_data),
                allowed_mentions=discord.AllowedMentions.none(),
            )
            # Store direct publications too, so the dashboard counts both direct and approved releases.
            with db_connect() as conn:
                cur = conn.execute(
                    "INSERT INTO releases(template_key,officer,payload_json,embed_json,status,requested_by,channel_id,message_id,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (self.template_key, str(self.values.get("officer", "See published embed")), json.dumps(self.values, ensure_ascii=False), json.dumps(self.embed_data, ensure_ascii=False), "published",
                     str(self.author_id), str(channel.id), str(message.id), utc_now(), utc_now()),
                )
                release_record_id = cur.lastrowid
            audit(self.author_id, "release.published", "release", release_record_id, {"channel_id": channel.id, "message_id": message.id})
        except discord.HTTPException:
            log.exception("Failed to publish IA release")
            await interaction.response.send_message(
                "Discord could not publish the release. Check the bot's channel permissions.", ephemeral=True
            )
            return
        for child in self.children:
            child.disabled = True
        await interaction.response.edit_message(content=f"Release published in {channel.mention}.", view=self)

    @discord.ui.button(label="Cancel", style=discord.ButtonStyle.secondary, row=0)
    async def cancel(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        for child in self.children:
            child.disabled = True
        await interaction.response.edit_message(content="Release cancelled.", view=self)


class ReleaseModal(discord.ui.Modal):
    def __init__(self, release_type: str, config: dict[str, Any]):
        super().__init__(title="IA Release — Officer Details", timeout=300)
        self.release_type = release_type
        self.config = config
        self.rank = discord.ui.TextInput(label="Officer rank", placeholder="e.g. OFC., SGT., LT.", max_length=40)
        self.manual_name = discord.ui.TextInput(
            label="Officer name (manual takes priority)",
            placeholder="Enter a name or leave blank to select a Discord user",
            max_length=100, required=False,
        )
        self.case_reference = discord.ui.TextInput(
            label="Case/reference (optional)", placeholder="Optional internal reference",
            max_length=100, required=False,
        )
        self.add_item(self.rank)
        self.add_item(self.manual_name)
        self.add_item(self.case_reference)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        try:
            rank = safe_text(str(self.rank.value), 40)
        except ValueError as exc:
            await interaction.response.send_message(str(exc), ephemeral=True)
            return
        manual_name = str(self.manual_name.value).strip() or None
        view = OfficerSelectionView(
            release_type=self.release_type,
            config=self.config,
            rank=rank,
            manual_name=manual_name,
            case_reference=str(self.case_reference.value).strip() or None,
            author_id=interaction.user.id,
        )
        if manual_name:
            message = "Your manual officer name will be used. Click **Continue with manual name**."
        else:
            message = "Select the officer from the menu below."
        await interaction.response.send_message(message, view=view, ephemeral=True)


class OfficerSelectionView(discord.ui.View):
    def __init__(self, release_type: str, config: dict[str, Any], rank: str,
                 manual_name: str | None, case_reference: str | None, author_id: int):
        super().__init__(timeout=300)
        self.release_type = release_type
        self.config = config
        self.rank = rank
        self.manual_name = manual_name
        self.case_reference = case_reference
        self.author_id = author_id
        self.add_item(OfficerSelect())
        self.add_item(ManualOfficerButton())

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who started this release can use these controls.", ephemeral=True)
            return False
        return True

    async def continue_with(self, interaction: discord.Interaction, user: discord.User | discord.Member | None) -> None:
        if user is not None and isinstance(user, discord.Member) and self.config.get("validate_rank_role", False):
            with db_connect() as conn:
                rank_row = conn.execute("SELECT discord_role_id FROM ranks WHERE name=? AND enabled=1", (self.rank,)).fetchone()
            required_role_id = str(rank_row["discord_role_id"] or "") if rank_row else ""
            if required_role_id and not any(str(role.id) == required_role_id for role in user.roles):
                await interaction.response.send_message("The selected officer does not have the Discord role linked to this rank.", ephemeral=True)
                return
        try:
            officer = display_officer(user, self.manual_name)
        except ValueError as exc:
            await interaction.response.send_message(str(exc), ephemeral=True)
            return
        if self.release_type == "investigation":
            await publish_preview(
                interaction, self.config, "investigation",
                {"rank": self.rank, "officer": officer, "case_reference": self.case_reference or ""},
                self.author_id,
            )
            return
        try:
            categories = load_violation_categories()
        except (OSError, ValueError, KeyError) as exc:
            await interaction.response.send_message(f"Could not load the Excel violation matrix: {exc}", ephemeral=True)
            return
        view = ViolationPickerView(self.config, self.rank, officer, self.case_reference, self.author_id, categories)
        await interaction.response.send_message(
            "Choose a category, select one or more offenses, then add them to your release.",
            view=view, ephemeral=True,
        )


class OfficerSelect(discord.ui.UserSelect):
    def __init__(self):
        super().__init__(placeholder="Select a Discord officer", min_values=1, max_values=1, row=0)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: OfficerSelectionView = self.view  # type: ignore[assignment]
        await view.continue_with(interaction, self.values[0] if self.values else None)


class ManualOfficerButton(discord.ui.Button):
    def __init__(self):
        super().__init__(label="Continue with manual name", style=discord.ButtonStyle.primary, row=1)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: OfficerSelectionView = self.view  # type: ignore[assignment]
        if not view.manual_name:
            await interaction.response.send_message(
                "No manual name was entered. Start again and enter a name, or select a Discord user.", ephemeral=True
            )
            return
        await view.continue_with(interaction, None)


class CategorySelect(discord.ui.Select):
    def __init__(self, categories: list[dict[str, Any]]):
        options = [discord.SelectOption(label=cat["name"][:100], value=cat["key"],
                                        description=f"{len(cat['records'])} offenses") for cat in categories[:25]]
        super().__init__(placeholder="1. Choose an offense category", min_values=1, max_values=1,
                         options=options, row=0)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: ViolationPickerView = self.view  # type: ignore[assignment]
        # Save any currently highlighted choices before switching categories.
        view.commit_current_selection()
        view.category_key = self.values[0]
        view.page = 0
        view.current_page_selection = []
        view.refresh_offense_select()
        await interaction.response.edit_message(content=view.status_text(), view=view)


class OffenseSelect(discord.ui.Select):
    def __init__(self, records: list[dict[str, str]], page: int):
        start = page * 25
        page_records = records[start:start + 25]
        options = []
        for record in page_records:
            description = " • ".join(x for x in (record.get("id", ""), record.get("severity", "")) if x)
            options.append(discord.SelectOption(
                label=record["offense"][:100], value=record["key"],
                description=description[:100] if description else None,
            ))
        super().__init__(placeholder=f"2. Select offenses (page {page + 1}) — multi-select",
                         min_values=1, max_values=max(1, len(options)), options=options, row=1)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: ViolationPickerView = self.view  # type: ignore[assignment]
        view.current_page_selection = list(self.values)
        await interaction.response.edit_message(content=view.status_text(), view=view)


class ViolationPickerView(discord.ui.View):
    def __init__(self, config: dict[str, Any], rank: str, officer: str, case_reference: str | None,
                 author_id: int, categories: list[dict[str, Any]]):
        super().__init__(timeout=600)
        self.config = config
        self.rank = rank
        self.officer = officer
        self.case_reference = case_reference
        self.author_id = author_id
        self.categories = categories
        self.records_by_key = all_violation_records(categories)
        self.category_key: str | None = None
        self.page = 0
        self.current_page_selection: list[str] = []
        self.selected_violations: dict[str, dict[str, str]] = {}
        self.add_item(CategorySelect(categories))

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who started this release can use these controls.", ephemeral=True)
            return False
        return True

    def current_category(self) -> dict[str, Any] | None:
        return next((cat for cat in self.categories if cat["key"] == self.category_key), None)

    def refresh_offense_select(self) -> None:
        for item in list(self.children):
            if isinstance(item, OffenseSelect):
                self.remove_item(item)
        category = self.current_category()
        if category:
            self.page = min(self.page, max(0, (len(category["records"]) - 1) // 25))
            self.add_item(OffenseSelect(category["records"], self.page))

    def commit_current_selection(self) -> None:
        for key in self.current_page_selection:
            if key in self.records_by_key:
                self.selected_violations[key] = self.records_by_key[key]
        self.current_page_selection = []

    def status_text(self) -> str:
        category = self.current_category()
        category_text = category["name"] if category else "not selected"
        selected_count = len(self.selected_violations) + len(set(self.current_page_selection) - set(self.selected_violations))
        selected_names = [r["offense"] for r in self.selected_violations.values()]
        if len(selected_names) > 5:
            selection_summary = ", ".join(selected_names[:5]) + f", and {len(selected_names) - 5} more"
        else:
            selection_summary = ", ".join(selected_names) if selected_names else "none added yet"
        return (f"**Category:** {category_text}\n**Selected offenses:** {selected_count}\n"
                f"**Added:** {selection_summary}\nSelect offenses, then click **Add selected**. You can switch categories and add more.")

    @discord.ui.button(label="Previous page", style=discord.ButtonStyle.secondary, row=2)
    async def previous_page(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        category = self.current_category()
        if not category or self.page <= 0:
            await interaction.response.send_message("There is no previous page.", ephemeral=True)
            return
        self.commit_current_selection()
        self.page -= 1
        self.refresh_offense_select()
        await interaction.response.edit_message(content=self.status_text(), view=self)

    @discord.ui.button(label="Add selected", style=discord.ButtonStyle.primary, row=2)
    async def add_selected(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        if not self.current_page_selection:
            await interaction.response.send_message("Select at least one offense first.", ephemeral=True)
            return
        self.commit_current_selection()
        await interaction.response.edit_message(content=self.status_text(), view=self)

    @discord.ui.button(label="Next page", style=discord.ButtonStyle.secondary, row=2)
    async def next_page(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        category = self.current_category()
        if not category or (self.page + 1) * 25 >= len(category["records"]):
            await interaction.response.send_message("There is no next page.", ephemeral=True)
            return
        self.commit_current_selection()
        self.page += 1
        self.refresh_offense_select()
        await interaction.response.edit_message(content=self.status_text(), view=self)

    @discord.ui.button(label="Continue to punishments", style=discord.ButtonStyle.success, row=3)
    async def continue_to_punishments(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        self.commit_current_selection()
        if not self.selected_violations:
            await interaction.response.send_message("Add at least one offense before continuing.", ephemeral=True)
            return
        view = PunishmentPickerView(
            self.config, self.rank, self.officer, self.case_reference, self.author_id,
            list(self.selected_violations.values()),
        )
        await interaction.response.edit_message(
            content=f"Selected {len(self.selected_violations)} offense(s). Choose one or more punishments.", view=view
        )


class PunishmentSelect(discord.ui.Select):
    def __init__(self, punishments: dict[str, Any]):
        options = []
        for key, item in list(punishments.items())[:25]:
            name = str(item.get("name", key)) if isinstance(item, dict) else str(item)
            options.append(discord.SelectOption(label=name[:100], value=key))
        super().__init__(placeholder="Select one or more punishments", min_values=1,
                         max_values=max(1, len(options)), options=options, row=0)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: PunishmentPickerView = self.view  # type: ignore[assignment]
        view.selected_punishments = list(self.values)
        await interaction.response.edit_message(content=view.status_text(), view=view)


class CustomPunishmentModal(discord.ui.Modal):
    def __init__(self, view: "PunishmentPickerView"):
        super().__init__(title="Add custom punishment", timeout=300)
        self.picker = view
        self.custom = discord.ui.TextInput(
            label="Custom punishment(s)", placeholder="One per line", style=discord.TextStyle.paragraph,
            max_length=500, required=True,
        )
        self.add_item(self.custom)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        custom_items = [line.strip() for line in str(self.custom.value).replace(";", "\n").splitlines() if line.strip()]
        if not custom_items:
            await interaction.response.send_message("Enter at least one custom punishment.", ephemeral=True)
            return
        await self.picker.finish(interaction, custom_items)


class PunishmentPickerView(discord.ui.View):
    def __init__(self, config: dict[str, Any], rank: str, officer: str, case_reference: str | None,
                 author_id: int, violations: list[dict[str, str]]):
        super().__init__(timeout=600)
        self.config = config
        self.rank = rank
        self.officer = officer
        self.case_reference = case_reference
        self.author_id = author_id
        self.violations = violations
        self.punishments = load_options("punishments.json")
        self.selected_punishments: list[str] = []
        self.add_item(PunishmentSelect(self.punishments))
        if not self.config.get("allow_custom_reasons_and_punishments", True):
            for child in self.children:
                if isinstance(child, discord.ui.Button) and child.label == "Add custom punishment":
                    child.disabled = True

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who started this release can use these controls.", ephemeral=True)
            return False
        return True

    def punishment_names(self) -> list[str]:
        names = []
        for key in self.selected_punishments:
            item = self.punishments.get(key, {})
            names.append(str(item.get("name", key)) if isinstance(item, dict) else str(item))
        return names

    def status_text(self) -> str:
        selected = self.punishment_names()
        return "**Selected punishments:** " + (", ".join(selected) if selected else "none") + "\nChoose punishments, then continue or add a custom punishment."

    async def finish(self, interaction: discord.Interaction, custom_items: list[str] | None = None) -> None:
        names = self.punishment_names()
        names.extend(custom_items or [])
        # Preserve order while removing exact duplicates.
        names = list(dict.fromkeys(name.strip() for name in names if name.strip()))
        if not names:
            await interaction.response.send_message("Select at least one punishment or add a custom one.", ephemeral=True)
            return
        reason_lines = []
        for record in self.violations:
            details = f"{record['id']} — {record['offense']}"
            if record.get("severity"):
                details += f" ({record['severity']})"
            reason_lines.append(details)
        await publish_preview(
            interaction, self.config, "disciplinary",
            {
                "rank": self.rank,
                "officer": self.officer,
                "reasons": format_list(reason_lines),
                "punishments": format_list(names),
                "case_reference": self.case_reference or "",
            },
            self.author_id,
        )

    @discord.ui.button(label="Continue", style=discord.ButtonStyle.success, row=1)
    async def continue_button(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await self.finish(interaction)

    @discord.ui.button(label="Add custom punishment", style=discord.ButtonStyle.secondary, row=1)
    async def custom_button(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(CustomPunishmentModal(self))


class ReleaseTypeView(discord.ui.View):
    def __init__(self, config: dict[str, Any], author_id: int):
        super().__init__(timeout=180)
        self.config = config
        self.author_id = author_id

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who started this command can select a release type.", ephemeral=True)
            return False
        if not authorized(interaction, self.config):
            await interaction.response.send_message("You do not have permission to create IA releases.", ephemeral=True)
            return False
        return True

    @discord.ui.button(label="Disciplinary Release", style=discord.ButtonStyle.primary)
    async def disciplinary(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(ReleaseModal("disciplinary", self.config))

    @discord.ui.button(label="Investigation Release", style=discord.ButtonStyle.primary)
    async def investigation(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(ReleaseModal("investigation", self.config))


# --------------------------- Integrated configuration panel ---------------------------
def can_manage_config(interaction: discord.Interaction, config: dict[str, Any]) -> bool:
    if interaction.guild is None or not isinstance(interaction.user, discord.Member):
        return False
    member = interaction.user
    if member.guild_permissions.administrator:
        return True
    role_ids = {int(value) for value in config.get("config_admin_role_ids", [])}
    return any(role.id in role_ids for role in member.roles)


def write_config_value(key: str, value: Any) -> dict[str, Any]:
    config = load_config()
    config[key] = value
    save_json("config.json", config)
    return config


class ConfigPanelView(discord.ui.View):
    def __init__(self, author_id: int):
        super().__init__(timeout=600)
        self.author_id = author_id

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        try:
            config = load_config()
        except (OSError, json.JSONDecodeError):
            config = {}
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who opened this panel can use it.", ephemeral=True)
            return False
        if not can_manage_config(interaction, config):
            await interaction.response.send_message("You do not have permission to manage the IA configuration.", ephemeral=True)
            return False
        return True

    def summary(self) -> str:
        config = load_config()
        templates = load_templates()
        channel_id = config.get("release_channel_id")
        channel_text = f"<#{channel_id}>" if channel_id else "Not configured"
        authorized_roles = config.get("authorized_role_ids", [])
        config_roles = config.get("config_admin_role_ids", [])
        return (
            "**RCPD Internal Affairs — Configuration Panel**\n"
            f"**Release channel:** {channel_text}\n"
            f"**IA release roles:** {len(authorized_roles)} configured\n"
            f"**Config/moderator roles:** {len(config_roles)} configured\n"
            f"**Templates:** {', '.join(str(v.get('name', k)) for k, v in templates.items())}\n\n"
            "Choose a setting below. Changes are saved to the `config` folder immediately."
        )

    @discord.ui.button(label="Release channel", style=discord.ButtonStyle.primary, row=0)
    async def channel_button(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.edit_message(content="Choose the channel where approved IA releases should be published.", view=ChannelConfigView(self.author_id))

    @discord.ui.button(label="IA release roles", style=discord.ButtonStyle.primary, row=0)
    async def release_roles_button(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.edit_message(content="Select the roles allowed to create and publish IA releases. This replaces the current list.", view=RoleConfigView(self.author_id, "authorized_role_ids", "Save IA roles"))

    @discord.ui.button(label="Config / moderator roles", style=discord.ButtonStyle.secondary, row=0)
    async def config_roles_button(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.edit_message(content="Select roles allowed to open and use this configuration panel. Server administrators always retain access.", view=RoleConfigView(self.author_id, "config_admin_role_ids", "Save config roles"))

    @discord.ui.button(label="Edit templates / embeds", style=discord.ButtonStyle.success, row=1)
    async def templates_button(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.edit_message(content="Choose which release embed to edit.", view=TemplateChooseView(self.author_id))

    @discord.ui.button(label="Other settings", style=discord.ButtonStyle.secondary, row=1)
    async def other_button(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(OtherSettingsModal())

    @discord.ui.button(label="Refresh panel", style=discord.ButtonStyle.secondary, row=1)
    async def refresh_button(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.edit_message(content=self.summary(), view=ConfigPanelView(self.author_id))


class ChannelConfigView(discord.ui.View):
    def __init__(self, author_id: int):
        super().__init__(timeout=300)
        self.author_id = author_id
        self.add_item(ReleaseChannelSelect())
        self.add_item(BackToConfigButton(author_id))

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who opened this panel can use it.", ephemeral=True)
            return False
        if not can_manage_config(interaction, load_config()):
            await interaction.response.send_message("You do not have permission to manage the configuration.", ephemeral=True)
            return False
        return True


class ReleaseChannelSelect(discord.ui.ChannelSelect):
    def __init__(self):
        super().__init__(placeholder="Select release channel", min_values=1, max_values=1,
                         channel_types=[discord.ChannelType.text, discord.ChannelType.news])

    async def callback(self, interaction: discord.Interaction) -> None:
        channel = self.values[0]
        write_config_value("release_channel_id", channel.id)
        await interaction.response.edit_message(content=f"Release channel saved: {channel.mention}", view=ConfigPanelView(interaction.user.id))


class RoleConfigView(discord.ui.View):
    def __init__(self, author_id: int, config_key: str, save_label: str):
        super().__init__(timeout=300)
        self.author_id = author_id
        self.config_key = config_key
        self.add_item(ConfigRoleSelect(config_key))
        self.add_item(BackToConfigButton(author_id))

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who opened this panel can use it.", ephemeral=True)
            return False
        if not can_manage_config(interaction, load_config()):
            await interaction.response.send_message("You do not have permission to manage the configuration.", ephemeral=True)
            return False
        return True


class ConfigRoleSelect(discord.ui.RoleSelect):
    def __init__(self, config_key: str):
        self.config_key = config_key
        super().__init__(placeholder="Select roles (up to 25)", min_values=0, max_values=25, row=0)

    async def callback(self, interaction: discord.Interaction) -> None:
        role_ids = [role.id for role in self.values]
        write_config_value(self.config_key, role_ids)
        label = "IA release roles" if self.config_key == "authorized_role_ids" else "Config/moderator roles"
        await interaction.response.edit_message(content=f"Saved {len(role_ids)} {label.lower()}.", view=ConfigPanelView(interaction.user.id))


class BackToConfigButton(discord.ui.Button):
    def __init__(self, author_id: int):
        super().__init__(label="Back to config panel", style=discord.ButtonStyle.secondary, row=1)
        self.author_id = author_id

    async def callback(self, interaction: discord.Interaction) -> None:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who opened this panel can use it.", ephemeral=True)
            return
        await interaction.response.edit_message(content=ConfigPanelView(self.author_id).summary(), view=ConfigPanelView(self.author_id))


class TemplateChooseView(discord.ui.View):
    def __init__(self, author_id: int, page: int = 0):
        super().__init__(timeout=300)
        self.author_id = author_id
        self.templates = list(load_templates().items())
        self.page = max(0, min(page, max(0, (len(self.templates) - 1) // 25)))
        self.add_item(TemplateSelect(self.templates, self.page))
        self.add_item(BackToConfigButton(author_id))

    @discord.ui.button(label="Previous templates", style=discord.ButtonStyle.secondary, row=1)
    async def previous_templates(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        if self.page <= 0:
            await interaction.response.send_message("Already on the first template page.", ephemeral=True)
            return
        await interaction.response.edit_message(content="Choose which release embed to edit.", view=TemplateChooseView(self.author_id, self.page - 1))

    @discord.ui.button(label="Next templates", style=discord.ButtonStyle.secondary, row=1)
    async def next_templates(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        if (self.page + 1) * 25 >= len(self.templates):
            await interaction.response.send_message("There are no more templates.", ephemeral=True)
            return
        await interaction.response.edit_message(content="Choose which release embed to edit.", view=TemplateChooseView(self.author_id, self.page + 1))

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who opened this panel can use it.", ephemeral=True)
            return False
        if not can_manage_config(interaction, load_config()):
            await interaction.response.send_message("You do not have permission to manage the configuration.", ephemeral=True)
            return False
        return True


class TemplateSelect(discord.ui.Select):
    def __init__(self, templates: list[tuple[str, dict[str, Any]]], page: int = 0):
        page_items = templates[page * 25:(page + 1) * 25]
        options = [discord.SelectOption(label=str(data.get("name", key))[:100], value=key) for key, data in page_items]
        super().__init__(placeholder=f"Select a release template (page {page + 1})", min_values=1, max_values=1, options=options, row=0)

    async def callback(self, interaction: discord.Interaction) -> None:
        await interaction.response.edit_message(content="Edit the selected embed's basic or advanced properties. Each save is written immediately.", view=TemplateEditorView(interaction.user.id, self.values[0]))


class TemplateEditorView(discord.ui.View):
    def __init__(self, author_id: int, template_key: str):
        super().__init__(timeout=600)
        self.author_id = author_id
        self.template_key = template_key

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who opened this editor can use it.", ephemeral=True)
            return False
        if not can_manage_config(interaction, load_config()):
            await interaction.response.send_message("You do not have permission to edit templates.", ephemeral=True)
            return False
        return True

    @discord.ui.button(label="Basic embed settings", style=discord.ButtonStyle.primary, row=0)
    async def basic(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(TemplateBasicModal(self.template_key))

    @discord.ui.button(label="Media / URLs / icons", style=discord.ButtonStyle.primary, row=0)
    async def advanced(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(TemplateAdvancedModal(self.template_key))

    @discord.ui.button(label="Edit embed fields", style=discord.ButtonStyle.primary, row=1)
    async def fields(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(TemplateFieldsModal(self.template_key))

    @discord.ui.button(label="Preview embed", style=discord.ButtonStyle.success, row=1)
    async def preview(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        templates = load_templates()
        data = templates.get(self.template_key, {})
        embed_data = data.get("embed", {})
        config = load_config()
        demo = {"rank": "SGT.", "officer": "Example Officer", "reasons": "EX-01 — Example offense", "punishments": "Official Warning", "case_reference": "CASE-001"}
        try:
            embed = build_release_embed(self.template_key, demo, config)
        except (ValueError, OSError, json.JSONDecodeError) as exc:
            await interaction.response.send_message(f"Template preview failed: {exc}", ephemeral=True)
            return
        await interaction.response.send_message(content=f"Preview: **{data.get('name', self.template_key)}**", embed=embed, ephemeral=True, allowed_mentions=discord.AllowedMentions.none())

    @discord.ui.button(label="Choose another template", style=discord.ButtonStyle.secondary, row=1)
    async def back(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.edit_message(content="Choose which release embed to edit.", view=TemplateChooseView(self.author_id))


class TemplateBasicModal(discord.ui.Modal):
    def __init__(self, template_key: str):
        super().__init__(title="Edit embed — basic settings", timeout=300)
        self.template_key = template_key
        templates = load_templates()
        item = templates[template_key]
        data = item.setdefault("embed", {})
        self.title_input = discord.ui.TextInput(label="Embed title", default=str(data.get("title", ""))[:256], max_length=256, required=False)
        self.description_input = discord.ui.TextInput(label="Description / body", default=str(data.get("description", item.get("template", "")))[:4000], style=discord.TextStyle.paragraph, max_length=4000, required=False)
        self.color_input = discord.ui.TextInput(label="Color hex, e.g. #202225", default=str(data.get("color", "#202225"))[:7], max_length=7, required=False)
        self.footer_input = discord.ui.TextInput(label="Footer text", default=str(data.get("footer", ""))[:2048], max_length=2048, required=False)
        self.author_input = discord.ui.TextInput(label="Author text", default=str(data.get("author", ""))[:256], max_length=256, required=False)
        for item_input in (self.title_input, self.description_input, self.color_input, self.footer_input, self.author_input):
            self.add_item(item_input)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        templates = load_templates()
        item = templates[self.template_key]
        data = item.setdefault("embed", {})
        color = str(self.color_input.value).strip()
        if color and not (len(color.lstrip("#")) in (3, 6) and all(c in "0123456789abcdefABCDEF" for c in color.lstrip("#"))):
            await interaction.response.send_message("Invalid color. Use a hex color such as #202225.", ephemeral=True)
            return
        data.update({"title": str(self.title_input.value), "description": str(self.description_input.value), "color": color or "#202225", "footer": str(self.footer_input.value), "author": str(self.author_input.value)})
        item["template"] = str(self.description_input.value)  # backwards compatibility
        save_json("templates.json", templates)
        await interaction.response.send_message("Basic embed settings saved. Open the config panel to continue editing.", ephemeral=True)


class TemplateAdvancedModal(discord.ui.Modal):
    def __init__(self, template_key: str):
        super().__init__(title="Edit embed — media and URLs", timeout=300)
        self.template_key = template_key
        templates = load_templates()
        item = templates[template_key]
        data = item.setdefault("embed", {})
        self.url_input = discord.ui.TextInput(label="Embed URL (optional)", default=str(data.get("url", ""))[:1000], max_length=1000, required=False)
        self.thumbnail_input = discord.ui.TextInput(label="Thumbnail image URL", default=str(data.get("thumbnail", ""))[:1000], max_length=1000, required=False)
        self.image_input = discord.ui.TextInput(label="Main image URL", default=str(data.get("image", ""))[:1000], max_length=1000, required=False)
        self.author_icon_input = discord.ui.TextInput(label="Author icon URL", default=str(data.get("author_icon_url", ""))[:1000], max_length=1000, required=False)
        self.footer_icon_input = discord.ui.TextInput(label="Footer icon URL", default=str(data.get("footer_icon_url", ""))[:1000], max_length=1000, required=False)
        for item_input in (self.url_input, self.thumbnail_input, self.image_input, self.author_icon_input, self.footer_icon_input):
            self.add_item(item_input)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        templates = load_templates()
        item = templates[self.template_key]
        data = item.setdefault("embed", {})
        data.update({
            "url": str(self.url_input.value).strip(),
            "thumbnail": str(self.thumbnail_input.value).strip(),
            "image": str(self.image_input.value).strip(),
            "author_icon_url": str(self.author_icon_input.value).strip(),
            "footer_icon_url": str(self.footer_icon_input.value).strip(),
        })
        save_json("templates.json", templates)
        await interaction.response.send_message("Embed media and URL settings saved.", ephemeral=True)


class TemplateFieldsModal(discord.ui.Modal):
    def __init__(self, template_key: str):
        super().__init__(title="Edit embed fields (JSON)", timeout=300)
        self.template_key = template_key
        templates = load_templates()
        data = templates[template_key].setdefault("embed", {})
        fields_json = json.dumps(data.get("fields", []), ensure_ascii=False)
        self.fields_input = discord.ui.TextInput(
            label="Embed fields as JSON list", default=fields_json[:4000],
            style=discord.TextStyle.paragraph, max_length=4000, required=False,
            placeholder='[{"name":"Case","value":"{case_reference}","inline":true}]',
        )
        self.add_item(self.fields_input)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        templates = load_templates()
        item = templates[self.template_key]
        data = item.setdefault("embed", {})
        raw_fields = str(self.fields_input.value).strip() or "[]"
        try:
            fields = json.loads(raw_fields)
        except json.JSONDecodeError as exc:
            await interaction.response.send_message(f"Fields JSON is invalid: {exc}", ephemeral=True)
            return
        if not isinstance(fields, list) or len(fields) > 25 or any(not isinstance(field, dict) for field in fields):
            await interaction.response.send_message("Fields must be a JSON list containing at most 25 objects with name, value, and optional inline.", ephemeral=True)
            return
        for field in fields:
            if not str(field.get("name", "")).strip() or "value" not in field:
                await interaction.response.send_message("Every field needs a non-empty 'name' and a 'value'.", ephemeral=True)
                return
        data["fields"] = fields
        save_json("templates.json", templates)
        await interaction.response.send_message("Embed fields saved.", ephemeral=True)


class OtherSettingsModal(discord.ui.Modal):
    def __init__(self):
        super().__init__(title="Other IA settings", timeout=300)
        config = load_config()
        self.footer = discord.ui.TextInput(label="Default embed footer", default=str(config.get("embed_footer", ""))[:2048], max_length=2048, required=False)
        self.max_length = discord.ui.TextInput(label="Max embed text length (1000–6000)", default=str(config.get("max_release_length", 4000)), max_length=4, required=True)
        self.allow_custom = discord.ui.TextInput(label="Allow custom punishments (true/false)", default=str(config.get("allow_custom_reasons_and_punishments", True)).lower(), max_length=5, required=True)
        self.add_item(self.footer)
        self.add_item(self.max_length)
        self.add_item(self.allow_custom)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        try:
            max_length = int(str(self.max_length.value).strip())
        except ValueError:
            await interaction.response.send_message("Maximum length must be a number.", ephemeral=True)
            return
        allow_custom = str(self.allow_custom.value).strip().lower()
        if not 1000 <= max_length <= 6000 or allow_custom not in {"true", "false"}:
            await interaction.response.send_message("Use a length from 1000 to 6000 and true/false for custom entries.", ephemeral=True)
            return
        config = load_config()
        config.update({"embed_footer": str(self.footer.value), "max_release_length": max_length, "allow_custom_reasons_and_punishments": allow_custom == "true"})
        save_json("config.json", config)
        await interaction.response.send_message("Other settings saved.", ephemeral=True)


class IABot(commands.Bot):
    async def setup_hook(self) -> None:
        if GUILD_ID and GUILD_ID.isdigit():
            guild = discord.Object(id=int(GUILD_ID))
            self.tree.copy_global_to(guild=guild)
            synced = await self.tree.sync(guild=guild)
            log.info("Synced %s command(s) to test guild %s", len(synced), GUILD_ID)
        else:
            synced = await self.tree.sync()
            log.info("Synced %s global command(s)", len(synced))


bot = IABot(command_prefix="!", intents=discord.Intents.default())


@bot.event
async def on_ready() -> None:
    log.info("Logged in as %s (ID: %s)", bot.user, bot.user.id if bot.user else "unknown")


@bot.tree.command(name="ia-release", description="Create an official RCPD Internal Affairs release.")
@app_commands.guild_only()
async def ia_release(interaction: discord.Interaction) -> None:
    try:
        config = load_config()
    except (OSError, json.JSONDecodeError) as exc:
        await interaction.response.send_message(f"Could not load bot configuration: {exc}", ephemeral=True)
        return
    if not authorized(interaction, config):
        await interaction.response.send_message("You do not have permission to create IA releases.", ephemeral=True)
        return
    await interaction.response.send_message(
        "Select the type of Internal Affairs release to create:",
        view=ReleaseTypeView(config, interaction.user.id), ephemeral=True,
    )


@bot.tree.command(name="ia-config", description="Open the integrated RCPD Internal Affairs configuration panel.")
@app_commands.guild_only()
async def ia_config(interaction: discord.Interaction) -> None:
    try:
        config = load_config()
    except (OSError, json.JSONDecodeError) as exc:
        await interaction.response.send_message(f"Could not load configuration: {exc}", ephemeral=True)
        return
    if not can_manage_config(interaction, config):
        await interaction.response.send_message("Only server administrators or configured config/moderator roles can use this panel.", ephemeral=True)
        return
    panel = ConfigPanelView(interaction.user.id)
    await interaction.response.send_message(panel.summary(), view=panel, ephemeral=True)


@bot.tree.command(name="ia-reload-config", description="Reload IA configuration and validate the Excel matrix.")
@app_commands.guild_only()
@app_commands.default_permissions(administrator=True)
async def ia_reload_config(interaction: discord.Interaction) -> None:
    if not isinstance(interaction.user, discord.Member) or not interaction.user.guild_permissions.administrator:
        await interaction.response.send_message("Only server administrators can reload the configuration.", ephemeral=True)
        return
    try:
        load_config()
        load_templates()
        load_options("punishments.json")
        categories = load_violation_categories()
    except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
        await interaction.response.send_message(f"Configuration error: {exc}", ephemeral=True)
        return
    offense_count = sum(len(category["records"]) for category in categories)
    await interaction.response.send_message(
        f"Configuration loaded. Found {len(categories)} offense categories and {offense_count} offenses in the Excel matrix.",
        ephemeral=True,
    )


# --------------------------- RCPD IA Add-on System ---------------------------
# SQLite-backed template versioning, ranks, approvals, cases, appeals, audit and backups.
import csv
import io
import re
import shutil
import sqlite3
import zipfile
from datetime import datetime, timezone

DB_PATH = ROOT / "data" / "ia.sqlite3"
BACKUP_DIR = ROOT / "backups"


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class ClosingConnection(sqlite3.Connection):
    def __exit__(self, exc_type, exc, tb):
        try:
            return super().__exit__(exc_type, exc, tb)
        finally:
            self.close()


def db_connect() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=15, factory=ClosingConnection)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def db_init() -> None:
    with db_connect() as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value_json TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS template_versions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            template_key TEXT NOT NULL,
            version INTEGER NOT NULL,
            name TEXT NOT NULL,
            template_json TEXT NOT NULL,
            actor_id TEXT NOT NULL,
            created_at TEXT NOT NULL,
            UNIQUE(template_key, version)
        );
        CREATE TABLE IF NOT EXISTS ranks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE COLLATE NOCASE,
            discord_role_id TEXT,
            sort_order INTEGER NOT NULL DEFAULT 100,
            enabled INTEGER NOT NULL DEFAULT 1,
            created_by TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS releases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            template_key TEXT NOT NULL,
            officer TEXT NOT NULL,
            payload_json TEXT NOT NULL,
            embed_json TEXT NOT NULL,
            status TEXT NOT NULL,
            requested_by TEXT NOT NULL,
            channel_id TEXT,
            message_id TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS approvals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            release_id INTEGER NOT NULL REFERENCES releases(id),
            status TEXT NOT NULL DEFAULT 'pending',
            requested_by TEXT NOT NULL,
            decided_by TEXT,
            decision_note TEXT,
            created_at TEXT NOT NULL,
            decided_at TEXT
        );
        CREATE TABLE IF NOT EXISTS cases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            case_code TEXT UNIQUE,
            title TEXT NOT NULL,
            officer TEXT NOT NULL,
            description TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL DEFAULT 'Open',
            assigned_to TEXT,
            created_by TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS case_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            case_id INTEGER NOT NULL REFERENCES cases(id) ON DELETE CASCADE,
            note TEXT NOT NULL,
            actor_id TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS appeals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            release_reference TEXT NOT NULL,
            officer TEXT NOT NULL,
            reason TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Pending',
            submitted_by TEXT NOT NULL,
            reviewer_id TEXT,
            decision_note TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS audit_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            actor_id TEXT NOT NULL,
            action TEXT NOT NULL,
            entity_type TEXT NOT NULL,
            entity_id TEXT,
            details_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_cases_status ON cases(status);
        CREATE INDEX IF NOT EXISTS idx_approvals_status ON approvals(status);
        CREATE INDEX IF NOT EXISTS idx_audit_created ON audit_logs(created_at);
        """)
        # Seed version history from the templates already installed with the bot.
        try:
            templates = load_templates()
        except Exception:
            templates = {}
        for key, template in templates.items():
            exists = conn.execute("SELECT 1 FROM template_versions WHERE template_key=? LIMIT 1", (key,)).fetchone()
            if not exists:
                conn.execute(
                    "INSERT INTO template_versions(template_key,version,name,template_json,actor_id,created_at) VALUES(?,?,?,?,?,?)",
                    (key, 1, str(template.get("name", key)), json.dumps(template, ensure_ascii=False), "system", utc_now()),
                )


def db_setting(key: str, default: Any = None) -> Any:
    with db_connect() as conn:
        row = conn.execute("SELECT value_json FROM settings WHERE key=?", (key,)).fetchone()
    return json.loads(row[0]) if row else default


def set_db_setting(key: str, value: Any) -> None:
    with db_connect() as conn:
        conn.execute("INSERT INTO settings(key,value_json,updated_at) VALUES(?,?,?) ON CONFLICT(key) DO UPDATE SET value_json=excluded.value_json, updated_at=excluded.updated_at",
                     (key, json.dumps(value, ensure_ascii=False), utc_now()))


def audit(actor_id: int | str, action: str, entity_type: str, entity_id: int | str | None = None, details: dict[str, Any] | None = None) -> None:
    with db_connect() as conn:
        conn.execute("INSERT INTO audit_logs(actor_id,action,entity_type,entity_id,details_json,created_at) VALUES(?,?,?,?,?,?)",
                     (str(actor_id), action, entity_type, str(entity_id) if entity_id is not None else None,
                      json.dumps(details or {}, ensure_ascii=False), utc_now()))


def is_module_enabled(name: str) -> bool:
    modules = db_setting("modules", {})
    return bool(modules.get(name, True))


def require_config_admin(interaction: discord.Interaction) -> bool:
    try:
        return can_manage_config(interaction, load_config())
    except Exception:
        return False


def is_approver(interaction: discord.Interaction) -> bool:
    if interaction.guild is None or not isinstance(interaction.user, discord.Member):
        return False
    if interaction.user.guild_permissions.administrator:
        return True
    config = load_config()
    allowed = {int(x) for x in config.get("approval_role_ids", [])}
    return any(role.id in allowed for role in interaction.user.roles)


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")
    return slug[:48] or "custom_template"


def record_template_version(template_key: str, template: dict[str, Any], actor_id: int | str) -> None:
    with db_connect() as conn:
        row = conn.execute("SELECT COALESCE(MAX(version),0) FROM template_versions WHERE template_key=?", (template_key,)).fetchone()
        version = int(row[0]) + 1
        conn.execute("INSERT INTO template_versions(template_key,version,name,template_json,actor_id,created_at) VALUES(?,?,?,?,?,?)",
                     (template_key, version, str(template.get("name", template_key)), json.dumps(template, ensure_ascii=False), str(actor_id), utc_now()))
    audit(actor_id, "template.version_created", "template", template_key, {"version": version})


class TemplateCreateModal(discord.ui.Modal, title="Create release template"):
    template_name = discord.ui.TextInput(label="Template name", max_length=80, required=True, placeholder="e.g. Reinstatement Release")
    description = discord.ui.TextInput(label="Embed description / body", style=discord.TextStyle.paragraph, max_length=4000, required=True,
                                       placeholder="Use {officer}, {rank}, {reasons}, {punishments}, {case_reference}")
    color = discord.ui.TextInput(label="Embed color hex", default="#202225", max_length=7, required=True)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        if not require_config_admin(interaction):
            await interaction.response.send_message("You cannot manage templates.", ephemeral=True)
            return
        if not is_module_enabled("templates"):
            await interaction.response.send_message("The Templates module is disabled.", ephemeral=True)
            return
        name = str(self.template_name.value).strip()
        body = str(self.description.value).strip()
        color = str(self.color.value).strip()
        if not body or not re.fullmatch(r"#?[0-9a-fA-F]{6}", color):
            await interaction.response.send_message("Description is required and color must be a hex value such as #202225.", ephemeral=True)
            return
        templates = load_templates()
        key = slugify(name)
        base, suffix, index = key, 2, 2
        while key in templates:
            key = f"{base}_{index}"
            index += 1
        item = {"name": name, "template": body, "embed": {"title": name, "description": body, "color": "#" + color.lstrip("#"),
                "footer": str(load_config().get("embed_footer", "RCPD Internal Affairs")), "author": "", "fields": []},
                "custom_template": True, "approval_required": False, "enabled": True}
        templates[key] = item
        save_json("templates.json", templates)
        record_template_version(key, item, interaction.user.id)
        await interaction.response.send_message(f"Created template **{name}** (`{key}`). It is available in `/ia-release` under Custom Templates.", ephemeral=True)


class RankCreateModal(discord.ui.Modal, title="Add rank"):
    rank_name = discord.ui.TextInput(label="Rank display name", max_length=80, required=True)
    role_id = discord.ui.TextInput(label="Optional linked Discord role ID", max_length=25, required=False, placeholder="Leave blank for no role link")
    sort_order = discord.ui.TextInput(label="Sort order (lower comes first)", default="100", max_length=5, required=True)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        if not require_config_admin(interaction):
            await interaction.response.send_message("You cannot manage ranks.", ephemeral=True)
            return
        name = str(self.rank_name.value).strip()
        role = str(self.role_id.value).strip()
        try:
            order = int(str(self.sort_order.value).strip())
            if role and not role.isdigit():
                raise ValueError
        except ValueError:
            await interaction.response.send_message("Role ID must be numeric and sort order must be an integer.", ephemeral=True)
            return
        try:
            with db_connect() as conn:
                conn.execute("INSERT INTO ranks(name,discord_role_id,sort_order,created_by,created_at) VALUES(?,?,?,?,?)",
                             (name, role or None, order, str(interaction.user.id), utc_now()))
        except sqlite3.IntegrityError:
            await interaction.response.send_message("That rank already exists.", ephemeral=True)
            return
        audit(interaction.user.id, "rank.created", "rank", name, {"role_id": role or None, "sort_order": order})
        await interaction.response.send_message(f"Rank **{name}** added.", ephemeral=True)


class RankAdminView(discord.ui.View):
    def __init__(self, author_id: int):
        super().__init__(timeout=300)
        self.author_id = author_id

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id or not require_config_admin(interaction):
            await interaction.response.send_message("You cannot manage ranks from this panel.", ephemeral=True)
            return False
        return True

    @discord.ui.button(label="Add rank", style=discord.ButtonStyle.success)
    async def add_rank(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(RankCreateModal())

    @discord.ui.button(label="List ranks", style=discord.ButtonStyle.primary)
    async def list_ranks(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        with db_connect() as conn:
            rows = conn.execute("SELECT name,discord_role_id,sort_order,enabled FROM ranks ORDER BY sort_order,name").fetchall()
        if not rows:
            text_out = "No ranks configured yet. Use **Add rank** first."
        else:
            text_out = "\n".join(f"• **{r['name']}** — order {r['sort_order']} — role `{r['discord_role_id'] or 'not linked'}` — {'enabled' if r['enabled'] else 'disabled'}" for r in rows[:40])
        await interaction.response.send_message(text_out, ephemeral=True)

    @discord.ui.button(label="Back to Config", style=discord.ButtonStyle.secondary)
    async def back(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        panel = ConfigPanelView(interaction.user.id)
        await interaction.response.edit_message(content=panel.summary(), view=panel)


class ApprovalRulesModal(discord.ui.Modal, title="Approval rules"):
    template_key = discord.ui.TextInput(label="Template key (or * for all)", max_length=80, required=True, placeholder="disciplinary or *")
    trigger_punishments = discord.ui.TextInput(label="Trigger punishments (comma-separated)", style=discord.TextStyle.paragraph, max_length=1000, required=False,
                                               placeholder="suspension, termination, blacklist")
    approval_roles = discord.ui.TextInput(label="Approver Discord role IDs (comma-separated)", style=discord.TextStyle.paragraph, max_length=1000, required=False)
    approval_channel = discord.ui.TextInput(label="Approval channel ID (optional)", max_length=25, required=False)
    require_for_template = discord.ui.TextInput(label="Always require approval? true/false", default="false", max_length=5, required=True)

    def __init__(self):
        super().__init__(timeout=300)
        config = load_config()
        role_ids = config.get("approval_role_ids", [])
        channel_id = config.get("approval_channel_id")
        self.approval_roles.default = ",".join(str(x) for x in role_ids)
        self.approval_channel.default = str(channel_id or "")

    async def on_submit(self, interaction: discord.Interaction) -> None:
        if not require_config_admin(interaction):
            await interaction.response.send_message("You cannot manage approval rules.", ephemeral=True)
            return
        key = str(self.template_key.value).strip()
        required_text = str(self.require_for_template.value).strip().lower()
        if required_text not in {"true", "false"}:
            await interaction.response.send_message("Use true or false for the template approval setting.", ephemeral=True)
            return
        triggers = [x.strip().casefold() for x in str(self.trigger_punishments.value).replace("\n", ",").split(",") if x.strip()]
        raw_roles = [x.strip() for x in str(self.approval_roles.value).replace("\n", ",").split(",") if x.strip()]
        approval_channel_id = str(self.approval_channel.value).strip()
        if approval_channel_id and not approval_channel_id.isdigit():
            await interaction.response.send_message("Approval channel ID must be numeric.", ephemeral=True)
            return
        if any(not x.isdigit() for x in raw_roles):
            await interaction.response.send_message("Approver role IDs must be numeric Discord role IDs.", ephemeral=True)
            return
        rules = db_setting("approval_rules", {})
        rules[key] = {"always": required_text == "true", "punishments": triggers}
        set_db_setting("approval_rules", rules)
        config = load_config()
        config["approval_role_ids"] = [int(x) for x in raw_roles]
        config["approval_channel_id"] = int(approval_channel_id) if approval_channel_id else None
        save_json("config.json", config)
        audit(interaction.user.id, "approval.rules_updated", "approval_rules", key, {"always": required_text == "true", "punishments": triggers, "role_count": len(raw_roles)})
        await interaction.response.send_message(f"Approval rules saved for `{key}`. Approver roles configured: {len(raw_roles)}.", ephemeral=True)


class CaseCreateModal(discord.ui.Modal, title="Create IA case"):
    title_input = discord.ui.TextInput(label="Case title", max_length=100, required=True)
    officer_input = discord.ui.TextInput(label="Officer name", max_length=100, required=True)
    description_input = discord.ui.TextInput(label="Case description", style=discord.TextStyle.paragraph, max_length=1500, required=False)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        if not authorized(interaction, load_config()):
            await interaction.response.send_message("You do not have permission to create cases.", ephemeral=True)
            return
        if not is_module_enabled("cases"):
            await interaction.response.send_message("Case Management is disabled in the module settings.", ephemeral=True)
            return
        now = utc_now()
        with db_connect() as conn:
            cur = conn.execute("INSERT INTO cases(title,officer,description,created_by,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                               (str(self.title_input.value).strip(), str(self.officer_input.value).strip(), str(self.description_input.value).strip(), str(interaction.user.id), now, now))
            case_id = cur.lastrowid
            case_code = f"IA-{datetime.now(timezone.utc).year}-{case_id:04d}"
            conn.execute("UPDATE cases SET case_code=? WHERE id=?", (case_code, case_id))
        audit(interaction.user.id, "case.created", "case", case_code, {"officer": str(self.officer_input.value).strip()})
        await interaction.response.send_message(f"Created case **{case_code}** for **{self.officer_input.value}**.", ephemeral=True)


class AppealCreateModal(discord.ui.Modal, title="Submit appeal"):
    release_reference = discord.ui.TextInput(label="Release or case reference", max_length=100, required=True)
    officer = discord.ui.TextInput(label="Officer name", max_length=100, required=True)
    reason = discord.ui.TextInput(label="Reason for appeal", style=discord.TextStyle.paragraph, max_length=2000, required=True)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        if not is_module_enabled("appeals"):
            await interaction.response.send_message("Appeals are disabled.", ephemeral=True)
            return
        now = utc_now()
        with db_connect() as conn:
            cur = conn.execute("INSERT INTO appeals(release_reference,officer,reason,submitted_by,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                               (str(self.release_reference.value).strip(), str(self.officer.value).strip(), str(self.reason.value).strip(), str(interaction.user.id), now, now))
            appeal_id = cur.lastrowid
        audit(interaction.user.id, "appeal.submitted", "appeal", appeal_id, {"reference": str(self.release_reference.value).strip()})
        await interaction.response.send_message(f"Appeal submitted. Reference: **AP-{appeal_id:04d}**. IA reviewers can process it using `/ia-appeal-decide`.", ephemeral=True)


class PanelAddonButton(discord.ui.Button):
    def __init__(self, label: str, action: str, row: int, style: discord.ButtonStyle = discord.ButtonStyle.secondary):
        super().__init__(label=label, style=style, row=row)
        self.action = action

    async def callback(self, interaction: discord.Interaction) -> None:
        if not require_config_admin(interaction):
            await interaction.response.send_message("You cannot use this configuration function.", ephemeral=True)
            return
        if self.action == "create_template":
            await interaction.response.send_modal(TemplateCreateModal())
        elif self.action == "ranks":
            await interaction.response.edit_message(content="**Rank Manager** — add and review configured ranks.", view=RankAdminView(interaction.user.id))
        elif self.action == "approvals":
            await interaction.response.send_modal(ApprovalRulesModal())
        elif self.action == "case":
            await interaction.response.send_modal(CaseCreateModal())
        elif self.action == "appeal":
            await interaction.response.send_modal(AppealCreateModal())
        elif self.action == "stats":
            await interaction.response.send_message(stats_text(), ephemeral=True)
        elif self.action == "audit":
            await interaction.response.send_message(audit_text(), ephemeral=True)
        elif self.action == "backup":
            try:
                path = create_backup()
            except Exception as exc:
                await interaction.response.send_message(f"Backup failed: {exc}", ephemeral=True)
                return
            await interaction.response.send_message("Backup created. Download and store it somewhere safe.", file=discord.File(path), ephemeral=True)
        elif self.action == "modules":
            await interaction.response.send_message(module_status_text(), ephemeral=True)


def _config_panel_init_with_addons(self: ConfigPanelView, author_id: int) -> None:
    _original_config_panel_init(self, author_id)
    self.add_item(PanelAddonButton("Create template", "create_template", row=2, style=discord.ButtonStyle.success))
    self.add_item(PanelAddonButton("Ranks", "ranks", row=2))
    self.add_item(PanelAddonButton("Approval rules", "approvals", row=2))
    self.add_item(PanelAddonButton("Create case", "case", row=3))
    self.add_item(PanelAddonButton("Submit appeal", "appeal", row=3))
    self.add_item(PanelAddonButton("Statistics", "stats", row=3))
    self.add_item(PanelAddonButton("Audit log", "audit", row=4))
    self.add_item(PanelAddonButton("Backup", "backup", row=4))
    self.add_item(PanelAddonButton("Modules", "modules", row=4))


_original_config_panel_init = ConfigPanelView.__init__
ConfigPanelView.__init__ = _config_panel_init_with_addons  # type: ignore[method-assign]


class RankSelect(discord.ui.Select):
    def __init__(self, rows: list[sqlite3.Row]):
        options = [discord.SelectOption(label=str(row["name"])[:100], value=str(row["id"]),
                    description=(f"Linked role {row['discord_role_id']}" if row["discord_role_id"] else "No linked Discord role")) for row in rows[:25]]
        super().__init__(placeholder="Select the officer's rank", min_values=1, max_values=1, options=options, row=0)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: RankSelectView = self.view  # type: ignore[assignment]
        row = next((r for r in view.rows if str(r["id"]) == self.values[0]), None)
        if row is None:
            await interaction.response.send_message("That rank no longer exists.", ephemeral=True)
            return
        view.rank_name = str(row["name"])
        view.rank_role_id = str(row["discord_role_id"] or "")
        if view.custom_template_key:
            await interaction.response.send_modal(CustomTemplateReleaseModal(view.custom_template_key, view.rank_name, view.config))
        else:
            next_view = OfficerSelectionView(view.release_type or "disciplinary", view.config, view.rank_name,
                view.manual_name, view.case_reference, view.author_id)
            await interaction.response.edit_message(content="Select the officer, or continue with the manual name if you entered one.", view=next_view)


class ManualRankModal(discord.ui.Modal, title="Enter custom rank"):
    rank_input = discord.ui.TextInput(label="Rank name", max_length=40, required=True)

    def __init__(self, parent: "RankSelectView"):
        super().__init__(timeout=300)
        self.parent_view = parent

    async def on_submit(self, interaction: discord.Interaction) -> None:
        rank = str(self.rank_input.value).strip()
        if not rank:
            await interaction.response.send_message("Rank cannot be empty.", ephemeral=True)
            return
        if self.parent_view.custom_template_key:
            next_view = ContinueCustomTemplateView(self.parent_view.custom_template_key, rank, self.parent_view.config, interaction.user.id)
            await interaction.response.send_message("Rank saved. Continue to the custom release details.", view=next_view, ephemeral=True)
        else:
            next_view = OfficerSelectionView(self.parent_view.release_type or "disciplinary", self.parent_view.config, rank,
                self.parent_view.manual_name, self.parent_view.case_reference, self.parent_view.author_id)
            await interaction.response.send_message("Select the officer, or continue with the manual name.", view=next_view, ephemeral=True)


class RankSelectView(discord.ui.View):
    def __init__(self, config: dict[str, Any], author_id: int, release_type: str | None = None,
                 manual_name: str | None = None, case_reference: str | None = None, custom_template_key: str | None = None):
        super().__init__(timeout=300)
        self.config = config
        self.author_id = author_id
        self.release_type = release_type
        self.manual_name = manual_name
        self.case_reference = case_reference
        self.custom_template_key = custom_template_key
        self.rank_name = ""
        self.rank_role_id = ""
        with db_connect() as conn:
            self.rows = conn.execute("SELECT id,name,discord_role_id,sort_order FROM ranks WHERE enabled=1 ORDER BY sort_order,name").fetchall()
        if self.rows:
            self.add_item(RankSelect(self.rows))
        self.add_item(ManualRankButton())

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("Only the person who started this release can use these controls.", ephemeral=True)
            return False
        return True


class ManualRankButton(discord.ui.Button):
    def __init__(self):
        super().__init__(label="Enter rank manually", style=discord.ButtonStyle.secondary, row=1)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: RankSelectView = self.view  # type: ignore[assignment]
        await interaction.response.send_modal(ManualRankModal(view))


# Replace the old rank text input with a selectable rank list while retaining a manual fallback.
_original_release_modal_init = ReleaseModal.__init__

def _release_modal_init_rank_select(self: ReleaseModal, release_type: str, config: dict[str, Any]):
    discord.ui.Modal.__init__(self, title="IA Release — Officer Details", timeout=300)
    self.release_type = release_type
    self.config = config
    self.manual_name = discord.ui.TextInput(label="Officer name (manual takes priority)", placeholder="Enter a name or leave blank to select a Discord user", max_length=100, required=False)
    self.case_reference = discord.ui.TextInput(label="Case/reference (optional)", placeholder="Optional internal reference", max_length=100, required=False)
    self.add_item(self.manual_name)
    self.add_item(self.case_reference)

async def _release_modal_submit_rank_select(self: ReleaseModal, interaction: discord.Interaction) -> None:
    manual_name = str(self.manual_name.value).strip() or None
    case_reference = str(self.case_reference.value).strip() or None
    view = RankSelectView(self.config, interaction.user.id, release_type=self.release_type,
                          manual_name=manual_name, case_reference=case_reference)
    text_out = "Choose a configured rank, or enter one manually."
    if not view.rows:
        text_out = "No ranks are configured yet. Enter a rank manually now, or add ranks in `/ia-config`."
    await interaction.response.send_message(text_out, view=view, ephemeral=True)

ReleaseModal.__init__ = _release_modal_init_rank_select  # type: ignore[method-assign]
ReleaseModal.on_submit = _release_modal_submit_rank_select  # type: ignore[method-assign]


class CustomTemplateSelect(discord.ui.Select):
    def __init__(self, templates: list[tuple[str, dict[str, Any]]], page: int = 0):
        page_items = templates[page * 25:(page + 1) * 25]
        options = [discord.SelectOption(label=str(value.get("name", key))[:100], value=key) for key, value in page_items]
        super().__init__(placeholder=f"Choose a custom template (page {page + 1})", min_values=1, max_values=1, options=options, row=0)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: CustomTemplateSelectView = self.view  # type: ignore[assignment]
        rank_view = RankSelectView(view.config, view.author_id, custom_template_key=self.values[0])
        await interaction.response.edit_message(content="Choose a rank for this custom template.", view=rank_view)


class CustomTemplateSelectView(discord.ui.View):
    def __init__(self, config: dict[str, Any], author_id: int, page: int = 0):
        super().__init__(timeout=300)
        self.config = config
        self.author_id = author_id
        self.templates = [(key, value) for key, value in load_templates().items() if value.get("custom_template") and value.get("enabled", True)]
        self.page = max(0, min(page, max(0, (len(self.templates) - 1) // 25)))
        self.add_item(CustomTemplateSelect(self.templates, self.page))

    @discord.ui.button(label="Previous", style=discord.ButtonStyle.secondary, row=1)
    async def previous_page(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        if self.page <= 0:
            await interaction.response.send_message("Already on the first page.", ephemeral=True)
            return
        await interaction.response.edit_message(content="Choose a custom release template.", view=CustomTemplateSelectView(self.config, self.author_id, self.page - 1))

    @discord.ui.button(label="Next", style=discord.ButtonStyle.secondary, row=1)
    async def next_page(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        if (self.page + 1) * 25 >= len(self.templates):
            await interaction.response.send_message("There are no more custom templates.", ephemeral=True)
            return
        await interaction.response.edit_message(content="Choose a custom release template.", view=CustomTemplateSelectView(self.config, self.author_id, self.page + 1))

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id or not authorized(interaction, self.config):
            await interaction.response.send_message("You cannot use this release selector.", ephemeral=True)
            return False
        return True


class CustomTemplateReleaseModal(discord.ui.Modal):
    def __init__(self, template_key: str, rank: str, config: dict[str, Any]):
        super().__init__(title="Custom template release", timeout=300)
        self.template_key = template_key
        self.rank = rank
        self.config = config
        self.officer = discord.ui.TextInput(label="Officer name", max_length=100, required=True)
        self.case_reference = discord.ui.TextInput(label="Case/reference (optional)", max_length=100, required=False)
        self.reasons = discord.ui.TextInput(label="Reasons / violations (optional)", style=discord.TextStyle.paragraph, max_length=1000, required=False)
        self.punishments = discord.ui.TextInput(label="Punishments / outcome (optional)", style=discord.TextStyle.paragraph, max_length=1000, required=False)
        for item in (self.officer, self.case_reference, self.reasons, self.punishments):
            self.add_item(item)

    async def on_submit(self, interaction: discord.Interaction) -> None:
        officer = discord.utils.escape_mentions(str(self.officer.value).strip())
        if not officer:
            await interaction.response.send_message("Officer name is required.", ephemeral=True)
            return
        values = {"officer": officer, "rank": self.rank, "case_reference": str(self.case_reference.value).strip(),
                  "reasons": str(self.reasons.value).strip() or "Not specified", "punishments": str(self.punishments.value).strip() or "Not specified"}
        await publish_preview(interaction, self.config, self.template_key, values, interaction.user.id)


class ContinueCustomTemplateView(discord.ui.View):
    def __init__(self, template_key: str, rank: str, config: dict[str, Any], author_id: int):
        super().__init__(timeout=300)
        self.template_key = template_key
        self.rank = rank
        self.config = config
        self.author_id = author_id

    @discord.ui.button(label="Continue to release details", style=discord.ButtonStyle.primary)
    async def continue_details(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        if interaction.user.id != self.author_id or not authorized(interaction, self.config):
            await interaction.response.send_message("You cannot continue this release.", ephemeral=True)
            return
        await interaction.response.send_modal(CustomTemplateReleaseModal(self.template_key, self.rank, self.config))


class CustomTemplatesButton(discord.ui.Button):
    def __init__(self):
        super().__init__(label="Custom templates", style=discord.ButtonStyle.success, row=1)

    async def callback(self, interaction: discord.Interaction) -> None:
        view: ReleaseTypeView = self.view  # type: ignore[assignment]
        if not is_module_enabled("templates"):
            await interaction.response.send_message("The Templates module is disabled.", ephemeral=True)
            return
        templates = [x for x in load_templates().values() if x.get("custom_template") and x.get("enabled", True)]
        if not templates:
            await interaction.response.send_message("No custom templates exist yet. Create one in `/ia-config`.", ephemeral=True)
            return
        await interaction.response.send_message("Choose a custom release template.", view=CustomTemplateSelectView(view.config, view.author_id), ephemeral=True)


_original_release_type_init = ReleaseTypeView.__init__
def _release_type_init_with_custom(self: ReleaseTypeView, config: dict[str, Any], author_id: int):
    _original_release_type_init(self, config, author_id)
    self.add_item(CustomTemplatesButton())
ReleaseTypeView.__init__ = _release_type_init_with_custom  # type: ignore[method-assign]


def approval_required(template_key: str, values: dict[str, str]) -> bool:
    templates = load_templates()
    item = templates.get(template_key, {})
    if bool(item.get("approval_required", False)):
        return True
    rules = db_setting("approval_rules", {})
    if rules.get(template_key, {}).get("always") or rules.get("*", {}).get("always"):
        return True
    triggers: set[str] = set()
    for key in (template_key, "*"):
        triggers.update(str(x).casefold() for x in rules.get(key, {}).get("punishments", []))
    punishment_text = str(values.get("punishments", "")).casefold()
    selected = {x.strip().casefold() for x in punishment_text.replace(";", ",").split(",") if x.strip()}
    selected.update(str(x).casefold() for x in load_options("punishments.json").keys() if x.casefold() in punishment_text)
    return bool(triggers.intersection(selected)) or any(trigger in punishment_text for trigger in triggers)


async def publish_approval_request(interaction: discord.Interaction, config: dict[str, Any], template_key: str,
                                   values: dict[str, str], author_id: int, embed: discord.Embed) -> None:
    now = utc_now()
    officer = str(values.get("officer", "Unknown"))
    with db_connect() as conn:
        cur = conn.execute("INSERT INTO releases(template_key,officer,payload_json,embed_json,status,requested_by,channel_id,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)",
                           (template_key, officer, json.dumps(values, ensure_ascii=False), json.dumps(embed.to_dict(), ensure_ascii=False), "pending_approval", str(author_id), str(config.get("release_channel_id") or ""), now, now))
        release_id = cur.lastrowid
        cur2 = conn.execute("INSERT INTO approvals(release_id,requested_by,created_at) VALUES(?,?,?)", (release_id, str(author_id), now))
        approval_id = cur2.lastrowid
    audit(author_id, "release.submitted_for_approval", "release", release_id, {"template": template_key, "approval_id": approval_id})
    approval_channel_id = config.get("approval_channel_id")
    channel = interaction.guild.get_channel(int(approval_channel_id)) if interaction.guild and approval_channel_id else None
    message = f"Release **#{release_id}** requires approval. Requested by <@{author_id}>. Use `/ia-approval-decide` if buttons are unavailable."
    if isinstance(channel, (discord.TextChannel, discord.Thread)):
        try:
            await channel.send(content=message, embed=embed, view=ApprovalReviewView(release_id), allowed_mentions=discord.AllowedMentions.none())
        except discord.HTTPException:
            log.exception("Could not post approval request")
    await interaction.response.send_message(f"This release requires approval and was queued as **IA Release #{release_id}**. It will not publish until approved.\n" + ("Approval request posted to the configured review channel." if isinstance(channel, (discord.TextChannel, discord.Thread)) else "No valid approval channel is configured; authorized reviewers can use `/ia-approvals` and `/ia-approval-decide` until one is configured."), ephemeral=True)


class ApprovalReviewView(discord.ui.View):
    def __init__(self, release_id: int):
        super().__init__(timeout=86400)
        self.release_id = release_id

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if not is_approver(interaction):
            await interaction.response.send_message("You are not configured as an IA approver.", ephemeral=True)
            return False
        with db_connect() as conn:
            row = conn.execute("SELECT status FROM releases WHERE id=?", (self.release_id,)).fetchone()
        if not row or row["status"] != "pending_approval":
            await interaction.response.send_message("This release is no longer awaiting approval.", ephemeral=True)
            return False
        return True

    @discord.ui.button(label="Approve and publish", style=discord.ButtonStyle.success)
    async def approve(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await approve_release(interaction, self.release_id, "Approved via review button")
        for child in self.children:
            child.disabled = True
        try:
            await interaction.message.edit(view=self)
        except discord.HTTPException:
            pass

    @discord.ui.button(label="Reject", style=discord.ButtonStyle.danger)
    async def reject(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        await interaction.response.send_modal(RejectApprovalModal(self.release_id))


class RejectApprovalModal(discord.ui.Modal, title="Reject IA release"):
    reason = discord.ui.TextInput(label="Reason for rejection", style=discord.TextStyle.paragraph, max_length=1000, required=True)

    def __init__(self, release_id: int):
        super().__init__(timeout=300)
        self.release_id = release_id

    async def on_submit(self, interaction: discord.Interaction) -> None:
        if not is_approver(interaction):
            await interaction.response.send_message("You are not configured as an IA approver.", ephemeral=True)
            return
        now = utc_now()
        with db_connect() as conn:
            row = conn.execute("SELECT status FROM releases WHERE id=?", (self.release_id,)).fetchone()
            if not row or row["status"] != "pending_approval":
                await interaction.response.send_message("This release is no longer pending.", ephemeral=True)
                return
            conn.execute("UPDATE releases SET status='rejected',updated_at=? WHERE id=?", (now, self.release_id))
            conn.execute("UPDATE approvals SET status='rejected',decided_by=?,decision_note=?,decided_at=? WHERE release_id=? AND status='pending'",
                         (str(interaction.user.id), str(self.reason.value).strip(), now, self.release_id))
        audit(interaction.user.id, "release.rejected", "release", self.release_id, {"reason": str(self.reason.value).strip()})
        await interaction.response.send_message(f"Release #{self.release_id} rejected and logged.", ephemeral=True)


async def approve_release(interaction: discord.Interaction, release_id: int, note: str = "Approved") -> None:
    if not is_approver(interaction):
        await interaction.response.send_message("You are not configured as an IA approver.", ephemeral=True)
        return
    now = utc_now()
    with db_connect() as conn:
        row = conn.execute("SELECT * FROM releases WHERE id=?", (release_id,)).fetchone()
        if not row or row["status"] != "pending_approval":
            await interaction.response.send_message("Release not found or no longer pending approval.", ephemeral=True)
            return
        config = load_config()
        channel_id = config.get("release_channel_id")
        channel = interaction.guild.get_channel(int(channel_id)) if interaction.guild and channel_id else None
        if not isinstance(channel, (discord.TextChannel, discord.Thread)):
            await interaction.response.send_message("The release channel is not configured.", ephemeral=True)
            return
        try:
            message = await channel.send(embed=discord.Embed.from_dict(json.loads(row["embed_json"])), allowed_mentions=discord.AllowedMentions.none())
        except discord.HTTPException:
            await interaction.response.send_message("Discord failed to publish the approved release. Check channel permissions.", ephemeral=True)
            return
        conn.execute("UPDATE releases SET status='published',message_id=?,updated_at=? WHERE id=?", (str(message.id), now, release_id))
        conn.execute("UPDATE approvals SET status='approved',decided_by=?,decision_note=?,decided_at=? WHERE release_id=? AND status='pending'",
                     (str(interaction.user.id), note, now, release_id))
    audit(interaction.user.id, "release.approved_and_published", "release", release_id, {"message_id": message.id})
    await interaction.response.send_message(f"Release #{release_id} approved and published in {channel.mention}.", ephemeral=True)


_BASE_PUBLISH_PREVIEW = publish_preview
async def publish_preview(interaction: discord.Interaction, config: dict[str, Any], release_type: str,
                          values: dict[str, str], author_id: int) -> None:
    if not authorized(interaction, config):
        await interaction.response.send_message("You do not have permission to create releases.", ephemeral=True)
        return
    try:
        embed = build_release_embed(release_type, values, config)
    except (ValueError, OSError, json.JSONDecodeError) as exc:
        await interaction.response.send_message(f"Template error: {exc}", ephemeral=True)
        return
    if approval_required(release_type, values):
        await publish_approval_request(interaction, config, release_type, values, author_id, embed)
        return
    # Direct publication preview remains the default for releases without an approval rule.
    max_length = int(config.get("max_release_length", 4000))
    length = len(embed.title or "") + len(embed.description or "") + sum(len(f.name) + len(f.value) for f in embed.fields)
    if length > min(max_length, 6000):
        await interaction.response.send_message(f"The rendered embed is too long ({length} characters).", ephemeral=True)
        return
    await interaction.response.send_message(content="**Preview — check the release before publishing:**", embed=embed,
        view=PublishView(embed.to_dict(), config, author_id, release_type, values), ephemeral=True, allowed_mentions=discord.AllowedMentions.none())


def stats_text() -> str:
    with db_connect() as conn:
        cases_total = conn.execute("SELECT COUNT(*) FROM cases").fetchone()[0]
        cases_open = conn.execute("SELECT COUNT(*) FROM cases WHERE status NOT IN ('Closed','Dismissed')").fetchone()[0]
        pending = conn.execute("SELECT COUNT(*) FROM approvals WHERE status='pending'").fetchone()[0]
        appeals_pending = conn.execute("SELECT COUNT(*) FROM appeals WHERE status='Pending'").fetchone()[0]
        releases = conn.execute("SELECT COUNT(*) FROM releases WHERE status='published'").fetchone()[0]
        ranks = conn.execute("SELECT COUNT(*) FROM ranks WHERE enabled=1").fetchone()[0]
    return ("**RCPD IA Statistics**\n"
            f"Published releases: **{releases}**\nCases: **{cases_total}** total / **{cases_open}** active\n"
            f"Pending approvals: **{pending}**\nPending appeals: **{appeals_pending}**\nActive ranks: **{ranks}**")


def audit_text(limit: int = 10) -> str:
    with db_connect() as conn:
        rows = conn.execute("SELECT actor_id,action,entity_type,entity_id,created_at FROM audit_logs ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
    if not rows:
        return "No audit events recorded yet."
    return "**Recent Audit Events**\n" + "\n".join(f"`{r['created_at'][:19]}` • <@{r['actor_id']}> • `{r['action']}` • {r['entity_type']} `{r['entity_id'] or '-'}`" for r in rows)


def module_status_text() -> str:
    modules = db_setting("modules", {"cases": True, "appeals": True, "templates": True})
    return "**Module status**\n" + "\n".join(f"• `{name}`: {'enabled' if enabled else 'disabled'}" for name, enabled in modules.items()) + "\nUse `/ia-module-toggle` to change a module."


def create_backup() -> Path:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    target = BACKUP_DIR / f"rcpd-ia-backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}.zip"
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in [DB_PATH, ROOT / "violations.xlsx", CONFIG_DIR / "config.json", CONFIG_DIR / "templates.json", CONFIG_DIR / "punishments.json"]:
            if path.exists():
                archive.write(path, arcname=str(path.relative_to(ROOT)))
    return target


# Slash commands for casework, appeals, approvals, modules, statistics and maintenance.
@bot.tree.command(name="ia-template-delete", description="Disable a custom IA template without deleting its history.")
@app_commands.guild_only()
@app_commands.describe(template_key="The template key shown in config/templates.json")
async def ia_template_delete(interaction: discord.Interaction, template_key: str) -> None:
    if not is_module_enabled("templates"):
        await interaction.response.send_message("The Templates module is disabled.", ephemeral=True)
        return
    if not require_config_admin(interaction):
        await interaction.response.send_message("Only config administrators can manage templates.", ephemeral=True)
        return
    templates = load_templates()
    if template_key not in templates or not templates[template_key].get("custom_template"):
        await interaction.response.send_message("Only custom templates can be disabled using this command.", ephemeral=True)
        return
    templates[template_key]["enabled"] = False
    save_json("templates.json", templates)
    record_template_version(template_key, templates[template_key], interaction.user.id)
    audit(interaction.user.id, "template.disabled", "template", template_key)
    await interaction.response.send_message(f"Template `{template_key}` disabled. Its history remains stored.", ephemeral=True)


@bot.tree.command(name="ia-approval-decide", description="Approve or reject a pending IA release by ID.")
@app_commands.guild_only()
@app_commands.describe(release_id="ID shown in the approval request", decision="approve or reject", reason="Required for rejection; optional note otherwise")
@app_commands.choices(decision=[app_commands.Choice(name="Approve", value="approve"), app_commands.Choice(name="Reject", value="reject")])
async def ia_approval_decide(interaction: discord.Interaction, release_id: int, decision: app_commands.Choice[str], reason: str = "") -> None:
    if not is_approver(interaction):
        await interaction.response.send_message("You are not configured as an IA approver.", ephemeral=True)
        return
    if decision.value == "approve":
        await approve_release(interaction, release_id, reason or "Approved via slash command")
        return
    if not reason.strip():
        await interaction.response.send_message("Provide a rejection reason.", ephemeral=True)
        return
    now = utc_now()
    with db_connect() as conn:
        row = conn.execute("SELECT status FROM releases WHERE id=?", (release_id,)).fetchone()
        if not row or row["status"] != "pending_approval":
            await interaction.response.send_message("Release not found or no longer pending.", ephemeral=True)
            return
        conn.execute("UPDATE releases SET status='rejected',updated_at=? WHERE id=?", (now, release_id))
        conn.execute("UPDATE approvals SET status='rejected',decided_by=?,decision_note=?,decided_at=? WHERE release_id=? AND status='pending'", (str(interaction.user.id), reason.strip(), now, release_id))
    audit(interaction.user.id, "release.rejected", "release", release_id, {"reason": reason.strip()})
    await interaction.response.send_message(f"Release #{release_id} rejected.", ephemeral=True)


@bot.tree.command(name="ia-approvals", description="List pending IA release approvals.")
@app_commands.guild_only()
async def ia_approvals(interaction: discord.Interaction) -> None:
    if not is_approver(interaction):
        await interaction.response.send_message("You are not configured as an IA approver.", ephemeral=True)
        return
    with db_connect() as conn:
        rows = conn.execute("SELECT id,template_key,officer,requested_by,created_at FROM releases WHERE status='pending_approval' ORDER BY id DESC LIMIT 15").fetchall()
    if not rows:
        text_out = "No pending approvals."
    else:
        text_out = "**Pending approvals**\n" + "\n".join(f"• Release `#{r['id']}` — `{r['template_key']}` — {r['officer']} — requested by <@{r['requested_by']}> — use `/ia-approval-decide release_id:{r['id']}`" for r in rows)
    await interaction.response.send_message(text_out, ephemeral=True)


@bot.tree.command(name="ia-case-create", description="Create an IA investigation case.")
@app_commands.guild_only()
@app_commands.describe(title="Short case title", officer="Officer involved", description="Case summary")
async def ia_case_create(interaction: discord.Interaction, title: str, officer: str, description: str = "") -> None:
    if not authorized(interaction, load_config()) or not is_module_enabled("cases"):
        await interaction.response.send_message("You do not have access to Case Management, or the module is disabled.", ephemeral=True)
        return
    now = utc_now()
    with db_connect() as conn:
        cur = conn.execute("INSERT INTO cases(title,officer,description,created_by,created_at,updated_at) VALUES(?,?,?,?,?,?)", (title[:100], officer[:100], description[:1500], str(interaction.user.id), now, now))
        cid = cur.lastrowid
        code = f"IA-{datetime.now(timezone.utc).year}-{cid:04d}"
        conn.execute("UPDATE cases SET case_code=? WHERE id=?", (code, cid))
    audit(interaction.user.id, "case.created", "case", code, {"officer": officer})
    await interaction.response.send_message(f"Created **{code}** — {title} (Officer: {officer}).", ephemeral=True)


@bot.tree.command(name="ia-cases", description="List recent IA cases.")
@app_commands.guild_only()
async def ia_cases(interaction: discord.Interaction) -> None:
    if not authorized(interaction, load_config()) or not is_module_enabled("cases"):
        await interaction.response.send_message("You do not have access to Case Management, or the module is disabled.", ephemeral=True)
        return
    with db_connect() as conn:
        rows = conn.execute("SELECT case_code,title,officer,status,assigned_to,updated_at FROM cases ORDER BY id DESC LIMIT 15").fetchall()
    text_out = "**Recent IA Cases**\n" + "\n".join(f"• `{r['case_code']}` — **{r['title']}** — {r['officer']} — `{r['status']}` — updated {r['updated_at'][:10]}" for r in rows) if rows else "No cases found."
    await interaction.response.send_message(text_out, ephemeral=True)


@bot.tree.command(name="ia-case-update", description="Update an IA case status or assign an investigator.")
@app_commands.guild_only()
@app_commands.describe(case_code="For example IA-2026-0001", status="New status", assigned_user="Optional investigator user ID")
@app_commands.choices(status=[app_commands.Choice(name=x, value=x) for x in ["Open", "Under Investigation", "Pending Approval", "Closed", "Dismissed", "Reopened"]])
async def ia_case_update(interaction: discord.Interaction, case_code: str, status: app_commands.Choice[str], assigned_user: str = "") -> None:
    if not authorized(interaction, load_config()) or not is_module_enabled("cases"):
        await interaction.response.send_message("You do not have access to Case Management, or the module is disabled.", ephemeral=True)
        return
    if assigned_user and not assigned_user.isdigit():
        await interaction.response.send_message("assigned_user must be a numeric Discord user ID.", ephemeral=True)
        return
    now = utc_now()
    with db_connect() as conn:
        row = conn.execute("SELECT id FROM cases WHERE case_code=?", (case_code.upper(),)).fetchone()
        if not row:
            await interaction.response.send_message("Case not found.", ephemeral=True)
            return
        conn.execute("UPDATE cases SET status=?,assigned_to=COALESCE(?,assigned_to),updated_at=? WHERE id=?", (status.value, assigned_user or None, now, row["id"]))
    audit(interaction.user.id, "case.updated", "case", case_code, {"status": status.value, "assigned_to": assigned_user or None})
    await interaction.response.send_message(f"Case `{case_code}` updated to **{status.value}**.", ephemeral=True)


@bot.tree.command(name="ia-case-note", description="Add an internal note to an IA case.")
@app_commands.guild_only()
async def ia_case_note(interaction: discord.Interaction, case_code: str, note: str) -> None:
    if not authorized(interaction, load_config()) or not is_module_enabled("cases"):
        await interaction.response.send_message("You do not have access to Case Management, or the module is disabled.", ephemeral=True)
        return
    with db_connect() as conn:
        row = conn.execute("SELECT id FROM cases WHERE case_code=?", (case_code.upper(),)).fetchone()
        if not row:
            await interaction.response.send_message("Case not found.", ephemeral=True)
            return
        conn.execute("INSERT INTO case_notes(case_id,note,actor_id,created_at) VALUES(?,?,?,?)", (row["id"], note[:2000], str(interaction.user.id), utc_now()))
        conn.execute("UPDATE cases SET updated_at=? WHERE id=?", (utc_now(), row["id"]))
    audit(interaction.user.id, "case.note_added", "case", case_code)
    await interaction.response.send_message(f"Note added to `{case_code}`.", ephemeral=True)


@bot.tree.command(name="ia-appeal", description="Submit an appeal for a release or case.")
@app_commands.guild_only()
async def ia_appeal(interaction: discord.Interaction, release_reference: str, officer: str, reason: str) -> None:
    if not is_module_enabled("appeals"):
        await interaction.response.send_message("Appeals module is disabled.", ephemeral=True)
        return
    now = utc_now()
    with db_connect() as conn:
        cur = conn.execute("INSERT INTO appeals(release_reference,officer,reason,submitted_by,created_at,updated_at) VALUES(?,?,?,?,?,?)", (release_reference[:100], officer[:100], reason[:2000], str(interaction.user.id), now, now))
        aid = cur.lastrowid
    audit(interaction.user.id, "appeal.submitted", "appeal", aid, {"reference": release_reference})
    await interaction.response.send_message(f"Appeal submitted as **AP-{aid:04d}**.", ephemeral=True)


@bot.tree.command(name="ia-appeals", description="List pending appeals for IA reviewers.")
@app_commands.guild_only()
async def ia_appeals(interaction: discord.Interaction) -> None:
    if not is_approver(interaction) or not is_module_enabled("appeals"):
        await interaction.response.send_message("Only configured IA approvers can review appeals.", ephemeral=True)
        return
    with db_connect() as conn:
        rows = conn.execute("SELECT id,release_reference,officer,reason,created_at FROM appeals WHERE status='Pending' ORDER BY id DESC LIMIT 10").fetchall()
    text_out = "**Pending appeals**\n" + "\n".join(f"• `AP-{r['id']:04d}` — `{r['release_reference']}` — {r['officer']} — {r['reason'][:150]} — decide with `/ia-appeal-decide appeal_id:{r['id']}`" for r in rows) if rows else "No pending appeals."
    await interaction.response.send_message(text_out, ephemeral=True)


@bot.tree.command(name="ia-appeal-decide", description="Approve or deny an appeal.")
@app_commands.guild_only()
@app_commands.choices(decision=[app_commands.Choice(name="Uphold", value="Uphold"), app_commands.Choice(name="Grant", value="Granted"), app_commands.Choice(name="Dismiss", value="Dismissed")])
async def ia_appeal_decide(interaction: discord.Interaction, appeal_id: int, decision: app_commands.Choice[str], note: str) -> None:
    if not is_approver(interaction) or not is_module_enabled("appeals"):
        await interaction.response.send_message("Only configured IA approvers can decide appeals.", ephemeral=True)
        return
    with db_connect() as conn:
        row = conn.execute("SELECT status FROM appeals WHERE id=?", (appeal_id,)).fetchone()
        if not row or row["status"] != "Pending":
            await interaction.response.send_message("Appeal not found or already decided.", ephemeral=True)
            return
        conn.execute("UPDATE appeals SET status=?,reviewer_id=?,decision_note=?,updated_at=? WHERE id=?", (decision.value, str(interaction.user.id), note[:1000], utc_now(), appeal_id))
    audit(interaction.user.id, "appeal.decided", "appeal", appeal_id, {"decision": decision.value, "note": note[:1000]})
    await interaction.response.send_message(f"Appeal `AP-{appeal_id:04d}` marked **{decision.value}**.", ephemeral=True)


@bot.tree.command(name="ia-stats", description="Show IA bot statistics.")
@app_commands.guild_only()
async def ia_stats(interaction: discord.Interaction) -> None:
    if not authorized(interaction, load_config()):
        await interaction.response.send_message("You do not have permission to view IA statistics.", ephemeral=True)
        return
    await interaction.response.send_message(stats_text(), ephemeral=True)


@bot.tree.command(name="ia-audit", description="Show recent IA configuration and case activity.")
@app_commands.guild_only()
async def ia_audit(interaction: discord.Interaction) -> None:
    if not require_config_admin(interaction) and not is_approver(interaction):
        await interaction.response.send_message("You do not have permission to view audit logs.", ephemeral=True)
        return
    await interaction.response.send_message(audit_text(15), ephemeral=True, allowed_mentions=discord.AllowedMentions.none())


@bot.tree.command(name="ia-backup", description="Create and download a backup of IA configuration and data (token excluded).")
@app_commands.guild_only()
async def ia_backup(interaction: discord.Interaction) -> None:
    if not require_config_admin(interaction):
        await interaction.response.send_message("Only config administrators can create backups.", ephemeral=True)
        return
    path = create_backup()
    audit(interaction.user.id, "backup.created", "backup", path.name)
    await interaction.response.send_message("Backup ready. Store it securely; it contains IA records but not the Discord token.", file=discord.File(path), ephemeral=True)


@bot.tree.command(name="ia-module-toggle", description="Enable or disable an IA add-on module.")
@app_commands.guild_only()
@app_commands.choices(module=[app_commands.Choice(name=x, value=x) for x in ["cases", "appeals", "templates"]])
async def ia_module_toggle(interaction: discord.Interaction, module: app_commands.Choice[str], enabled: bool) -> None:
    if not require_config_admin(interaction):
        await interaction.response.send_message("Only config administrators can change module settings.", ephemeral=True)
        return
    modules = db_setting("modules", {"cases": True, "appeals": True, "templates": True})
    modules[module.value] = enabled
    set_db_setting("modules", modules)
    audit(interaction.user.id, "module.toggled", "module", module.value, {"enabled": enabled})
    await interaction.response.send_message(f"Module `{module.value}` is now {'enabled' if enabled else 'disabled'}.", ephemeral=True)



@bot.tree.command(name="ia-template-duplicate", description="Duplicate an existing template under a new name.")
@app_commands.guild_only()
async def ia_template_duplicate(interaction: discord.Interaction, source_key: str, new_name: str) -> None:
    if not is_module_enabled("templates"):
        await interaction.response.send_message("The Templates module is disabled.", ephemeral=True)
        return
    if not require_config_admin(interaction):
        await interaction.response.send_message("Only config administrators can manage templates.", ephemeral=True)
        return
    templates = load_templates()
    if source_key not in templates:
        await interaction.response.send_message("Source template key was not found.", ephemeral=True)
        return
    new_name = new_name.strip()
    if not new_name:
        await interaction.response.send_message("New template name cannot be empty.", ephemeral=True)
        return
    key = slugify(new_name)
    base = key
    counter = 2
    while key in templates:
        key = f"{base}_{counter}"
        counter += 1
    clone = json.loads(json.dumps(templates[source_key]))
    clone["name"] = new_name[:80]
    clone["custom_template"] = True
    clone["enabled"] = True
    if "embed" in clone:
        clone["embed"]["title"] = new_name[:256]
    templates[key] = clone
    save_json("templates.json", templates)
    record_template_version(key, clone, interaction.user.id)
    await interaction.response.send_message(f"Duplicated `{source_key}` as **{new_name}** (`{key}`).", ephemeral=True)


@bot.tree.command(name="ia-rank-role-check", description="Require selected Discord officers to have their linked rank role.")
@app_commands.guild_only()
async def ia_rank_role_check(interaction: discord.Interaction, enabled: bool) -> None:
    if not require_config_admin(interaction):
        await interaction.response.send_message("Only config administrators can change rank validation.", ephemeral=True)
        return
    config = load_config()
    config["validate_rank_role"] = enabled
    save_json("config.json", config)
    audit(interaction.user.id, "rank.role_validation_changed", "config", "validate_rank_role", {"enabled": enabled})
    await interaction.response.send_message(f"Rank-role validation is now {'enabled' if enabled else 'disabled'}.", ephemeral=True)


@bot.tree.command(name="ia-officer-history", description="View IA case, release, and appeal history for an officer.")
@app_commands.guild_only()
async def ia_officer_history(interaction: discord.Interaction, officer_name: str) -> None:
    if not authorized(interaction, load_config()):
        await interaction.response.send_message("You do not have permission to view officer history.", ephemeral=True)
        return
    term = f"%{officer_name.strip()}%"
    with db_connect() as conn:
        cases_rows = conn.execute("SELECT case_code,title,status,updated_at FROM cases WHERE officer LIKE ? ORDER BY id DESC LIMIT 8", (term,)).fetchall()
        release_rows = conn.execute("SELECT id,template_key,status,created_at FROM releases WHERE officer LIKE ? ORDER BY id DESC LIMIT 8", (term,)).fetchall()
        appeal_rows = conn.execute("SELECT id,release_reference,status,created_at FROM appeals WHERE officer LIKE ? ORDER BY id DESC LIMIT 8", (term,)).fetchall()
    parts = [f"**IA history for:** {discord.utils.escape_mentions(officer_name)}"]
    parts.append("**Cases**\n" + ("\n".join(f"• `{r['case_code']}` — {r['title']} — `{r['status']}`" for r in cases_rows) if cases_rows else "No matching cases."))
    parts.append("**Releases**\n" + ("\n".join(f"• Release `#{r['id']}` — `{r['template_key']}` — `{r['status']}`" for r in release_rows) if release_rows else "No matching releases."))
    parts.append("**Appeals**\n" + ("\n".join(f"• `AP-{r['id']:04d}` — `{r['release_reference']}` — `{r['status']}`" for r in appeal_rows) if appeal_rows else "No matching appeals."))
    audit(interaction.user.id, "officer.history_viewed", "officer", officer_name)
    await interaction.response.send_message("\n\n".join(parts)[:1900], ephemeral=True, allowed_mentions=discord.AllowedMentions.none())

class RestoreConfirmView(discord.ui.View):
    def __init__(self, author_id: int, archive_bytes: bytes):
        super().__init__(timeout=180)
        self.author_id = author_id
        self.archive_bytes = archive_bytes

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id or not require_config_admin(interaction):
            await interaction.response.send_message("Only the authorized administrator who started the restore can confirm it.", ephemeral=True)
            return False
        return True

    @discord.ui.button(label="Confirm restore", style=discord.ButtonStyle.danger)
    async def confirm_restore(self, interaction: discord.Interaction, button: discord.ui.Button) -> None:
        allowed = {"data/ia.sqlite3", "config/config.json", "config/templates.json", "config/punishments.json", "violations.xlsx"}
        if len(self.archive_bytes) > 25 * 1024 * 1024:
            await interaction.response.send_message("Backup is too large to restore through Discord (limit 25 MB).", ephemeral=True)
            return
        try:
            with zipfile.ZipFile(io.BytesIO(self.archive_bytes), "r") as archive:
                names = {name for name in archive.namelist() if not name.endswith("/")}
                if not names.issubset(allowed):
                    await interaction.response.send_message("Backup contains unexpected files; restore cancelled.", ephemeral=True)
                    return
                required = {"data/ia.sqlite3", "config/config.json", "config/templates.json", "config/punishments.json"}
                if not required.issubset(names):
                    await interaction.response.send_message("Backup is incomplete. It must contain the database and all core configuration files.", ephemeral=True)
                    return
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    await interaction.response.send_message("Backup expands beyond the safe size limit; restore cancelled.", ephemeral=True)
                    return
                payloads = {name: archive.read(name) for name in names}
            # Validate JSON before touching active files.
            for name in ("config/config.json", "config/templates.json", "config/punishments.json"):
                json.loads(payloads[name].decode("utf-8"))
            ROOT.joinpath("data").mkdir(parents=True, exist_ok=True)
            staging = ROOT / "data" / "restore-staging"
            staging.mkdir(parents=True, exist_ok=True)
            staged_paths: dict[str, Path] = {}
            for name, raw in payloads.items():
                staged = staging / Path(name).name
                staged.write_bytes(raw)
                staged_paths[name] = staged
            test_conn = sqlite3.connect(staged_paths["data/ia.sqlite3"])
            try:
                integrity = test_conn.execute("PRAGMA integrity_check").fetchone()[0]
            finally:
                test_conn.close()
            if integrity != "ok":
                await interaction.response.send_message("Database integrity check failed. Active files were not changed.", ephemeral=True)
                return
            safety_backup = create_backup()
            for name, staged in staged_paths.items():
                target = ROOT / name
                target.parent.mkdir(parents=True, exist_ok=True)
                temporary = target.with_name(target.name + ".restore.tmp")
                shutil.copy2(staged, temporary)
                temporary.replace(target)
            audit(interaction.user.id, "backup.restored", "backup", safety_backup.name, {"source": "uploaded zip"})
            for child in self.children:
                child.disabled = True
            await interaction.response.edit_message(content=f"Restore completed. A safety backup was created as `{safety_backup.name}`. Restart the bot before further use to ensure all state is reloaded.", view=self)
        except (OSError, ValueError, KeyError, zipfile.BadZipFile, sqlite3.DatabaseError, json.JSONDecodeError) as exc:
            log.exception("Backup restore failed")
            await interaction.response.send_message(f"Restore failed; check the backup file. Details: {str(exc)[:300]}", ephemeral=True)


@bot.tree.command(name="ia-restore", description="Restore IA data from a downloaded backup ZIP; creates a safety backup first.")
@app_commands.guild_only()
async def ia_restore(interaction: discord.Interaction, backup: discord.Attachment) -> None:
    if not require_config_admin(interaction):
        await interaction.response.send_message("Only config administrators can restore backups.", ephemeral=True)
        return
    if not backup.filename.lower().endswith(".zip"):
        await interaction.response.send_message("Upload a backup ZIP created by `/ia-backup`.", ephemeral=True)
        return
    if backup.size > 25 * 1024 * 1024:
        await interaction.response.send_message("Backup is too large to restore through Discord (limit 25 MB).", ephemeral=True)
        return
    try:
        data = await backup.read()
        with zipfile.ZipFile(io.BytesIO(data), "r") as archive:
            names = {name for name in archive.namelist() if not name.endswith("/")}
            required = {"data/ia.sqlite3", "config/config.json", "config/templates.json", "config/punishments.json"}
            if not required.issubset(names):
                await interaction.response.send_message("This does not look like a complete RCPD IA backup ZIP.", ephemeral=True)
                return
    except (OSError, zipfile.BadZipFile) as exc:
        await interaction.response.send_message(f"Could not read backup ZIP: {exc}", ephemeral=True)
        return
    await interaction.response.send_message(
        "**Confirm restore?** This replaces the active IA database and core configuration with the uploaded backup. A safety backup of the current state will be created first. The bot should be restarted afterwards.",
        view=RestoreConfirmView(interaction.user.id, data), ephemeral=True,
    )


@bot.tree.command(name="ia-case-view", description="View a case summary and its internal notes.")
@app_commands.guild_only()
async def ia_case_view(interaction: discord.Interaction, case_code: str) -> None:
    if not authorized(interaction, load_config()) or not is_module_enabled("cases"):
        await interaction.response.send_message("You do not have access to Case Management, or the module is disabled.", ephemeral=True)
        return
    with db_connect() as conn:
        row = conn.execute("SELECT * FROM cases WHERE case_code=?", (case_code.upper(),)).fetchone()
        if not row:
            await interaction.response.send_message("Case not found.", ephemeral=True)
            return
        notes = conn.execute("SELECT note,actor_id,created_at FROM case_notes WHERE case_id=? ORDER BY id DESC LIMIT 10", (row["id"],)).fetchall()
    note_text = "\n".join(f"• `{note['created_at'][:10]}` — <@{note['actor_id']}>: {note['note'][:250]}" for note in notes) if notes else "No notes recorded."
    text_out = (f"**{row['case_code']} — {row['title']}**\nOfficer: {row['officer']}\nStatus: `{row['status']}`\n"
                f"Assigned to: {('<@' + row['assigned_to'] + '>') if row['assigned_to'] else 'Unassigned'}\n"
                f"Description: {row['description'] or '—'}\n\n**Recent notes**\n{note_text}")
    await interaction.response.send_message(text_out[:1900], ephemeral=True, allowed_mentions=discord.AllowedMentions.none())



# Keep database migrations and template version history ready before the bot logs in.
db_init()


if __name__ == "__main__":
    if not TOKEN or TOKEN == "put_your_bot_token_here":
        raise SystemExit("DISCORD_TOKEN is missing. Copy .env.example to .env and add your bot token.")
    bot.run(TOKEN)
